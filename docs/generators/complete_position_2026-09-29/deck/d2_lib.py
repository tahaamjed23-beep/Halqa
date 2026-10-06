# -*- coding: utf-8 -*-
"""
Halqa: the complete position, second edition (29 September 2026). Shared helpers and slide chrome.

House look: white pages, Georgia titles in lime 700, Arial body, hairline ledger rules instead of tinted cards,
native charts for numbers, exhibit captions and a source line on every page that carries a figure. Lime ramp only.
"""
import math, os, re
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK, XL_TICK_LABEL_POSITION
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, 'lockup-520.png')
MARK = os.path.join(HERE, 'mark-160.png')
MQ_LOGO = os.path.join(HERE, '..', 'logos', 'mashreq-logo-colour.png')
OUT = os.path.join(HERE, 'Halqa-Complete-Position.pptx')
DATE = '29 September 2026'


def C(h):
    return RGBColor.from_string(h)


LIME, L700, L600, L800 = C('6DC72A'), C('41801A'), C('57A81E'), C('2F6313')
LIME_LT, LIME_XLT, LIME_MID = C('E3F8CA'), C('F3FCE7'), C('A9E76A')
INK, INK2, GREY, GREY_LT, RULE, WHITE = C('0C1408'), C('26301F'), C('5F6B58'), C('A9B4A3'), C('CFD8C6'), C('FFFFFF')
PAPER = C('F6F7F4')
RED, AMBER, MQ = C('C0392B'), C('D9A21B'), C('D8602A')
DISPLAY, BODY = 'Georgia', 'Arial'
L, R, CEN = PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.CENTER
TOP, MIDDLE, BOTTOM = MSO_ANCHOR.TOP, MSO_ANCHOR.MIDDLE, MSO_ANCHOR.BOTTOM
W, H, M = 13.333, 7.5, 0.55
CW = W - 2 * M
Y0 = 1.62          # top of the content area
YS = 6.78          # source line


def E(v):
    return int(round(v * 914400))


prs = Presentation()
prs.slide_width, prs.slide_height = E(W), E(H)
BLANK = [lay for lay in prs.slide_layouts if lay.name == 'Blank'][0]
for lay in list(prs.slide_layouts):
    if lay.name != 'Blank':
        prs.slide_layouts.remove(lay)
PAGE = [0]
SECTION = ['']
TOC = []           # (page, title, section) for the contents and the dividers
EXH = [0]


# ------------------------------------------------------------------- text ---
def _runs(p, text, st):
    segs = []
    if isinstance(text, list):
        segs = [(t, dict(st, **s2)) for t, s2 in text]
    else:
        for k, seg in enumerate(text.split('**')):
            if seg:
                segs.append((seg, dict(st, bold=st.get('bold') or k % 2 == 1)))
    for seg, s in segs:
        r = p.add_run()
        r.text = seg.upper() if s.get('caps') else seg
        f = r.font
        f.size = Pt(s.get('size', 11))
        f.name = s.get('font', BODY)
        f.color.rgb = s.get('color', INK)
        if s.get('bold'):
            f.bold = True
        if s.get('italic'):
            f.italic = True
        if s.get('track'):
            r.font._rPr.set('spc', str(int(s['track'] * 100)))


def fill(tf, paras, **st):
    tf.word_wrap = True
    if isinstance(paras, str):
        paras = [paras]
    for i, item in enumerate(paras):
        text, own = (item if isinstance(item, tuple) else (item, {}))
        s = dict(st, **own)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if s.get('align', L) != L:
            p.alignment = s['align']
        elif p._p.pPr is not None and p._p.pPr.get('algn') is not None:
            del p._p.pPr.attrib['algn']
        p.line_spacing = s.get('spacing', 1.05)
        if s.get('gap') is not None:
            p.space_after = Pt(s['gap'])
        if s.get('before'):
            p.space_before = Pt(s['before'])
        if s.get('bullet') or s.get('num') or s.get('dash_ind'):
            pPr = p._p.get_or_add_pPr()
            ind = E(s.get('ind', 0.2))
            pPr.set('marL', str(ind))
            pPr.set('indent', str(-ind))
            if s.get('bullet') or s.get('num'):
                bc = etree.SubElement(pPr, qn('a:buClr'))
                etree.SubElement(bc, qn('a:srgbClr')).set('val', str(s.get('bcolor', LIME)))
            if s.get('bullet'):
                etree.SubElement(pPr, qn('a:buSzPct')).set('val', '70000')
                etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
                etree.SubElement(pPr, qn('a:buChar')).set('char', u'■')
            elif s.get('num'):
                etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
                etree.SubElement(pPr, qn('a:buAutoNum')).set('type', 'arabicPeriod')
        _runs(p, text, s)


def _frame(tf, pad, anchor, force=False):
    if isinstance(pad, (int, float)):
        pad = (pad, pad * 0.75, pad, pad * 0.75)
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [E(v) for v in pad]
    if anchor != TOP or force:
        tf.vertical_anchor = anchor


def tb(sl, x, y, w, h, paras, pad=0, anchor=TOP, **st):
    box = sl.shapes.add_textbox(E(x), E(y), E(w), E(h))
    _frame(box.text_frame, pad, anchor)
    fill(box.text_frame, paras, **st)
    return box


def _nostyle(shape_):
    st = shape_._element.find(qn('p:style'))
    if st is not None:
        shape_._element.remove(st)


def shape(sl, kind, x, y, w, h, color=None, line=None, lw=0.75, rot=0, paras=None, pad=0.12, anchor=TOP, **st):
    s = sl.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
    _nostyle(s)
    if color is not None:
        s.fill.solid()
        s.fill.fore_color.rgb = color
    else:
        s.fill.background()
    if line is not None:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    if rot:
        s.rotation = rot
    if paras is not None:
        _frame(s.text_frame, pad, anchor, force=True)
        fill(s.text_frame, paras, **st)
    return s


def rect(sl, x, y, w, h, color=None, **kw):
    return shape(sl, MSO_SHAPE.RECTANGLE, x, y, w, h, color, **kw)


def box(sl, x, y, w, h, paras, color=WHITE, line=RULE, lw=0.75, pad=0.12, anchor=TOP, **st):
    """Outlined diagram box, square corners."""
    return rect(sl, x, y, w, h, color, line=line, lw=lw, paras=paras, pad=pad, anchor=anchor, **st)


def oval(sl, x, y, d, color, **kw):
    return shape(sl, MSO_SHAPE.OVAL, x, y, d, d, color, **kw)


def line(sl, x1, y1, x2, y2, color=L700, lw=1.0, arrow=False, dash=False, head=False):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    _nostyle(c)
    c.line.color.rgb = color
    c.line.width = Pt(lw)
    ln = c.line._get_or_add_ln()
    if dash:
        etree.SubElement(ln, qn('a:prstDash')).set('val', 'dash')
    if head:
        he = etree.SubElement(ln, qn('a:headEnd'))
        he.set('type', 'triangle'); he.set('w', 'med'); he.set('len', 'med')
    if arrow:
        te = etree.SubElement(ln, qn('a:tailEnd'))
        te.set('type', 'triangle'); te.set('w', 'med'); te.set('len', 'med')
    return c


def rule(sl, x, y, w, color=RULE, lw=0.75):
    return line(sl, x, y, x + w, y, color, lw)


def vrule(sl, x, y, h, color=RULE, lw=0.75):
    return line(sl, x, y, x, y + h, color, lw)


def elbow(sl, pts, color=L700, lw=1.2, arrow=True):
    for i in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        line(sl, x1, y1, x2, y2, color, lw, arrow=arrow and i == len(pts) - 2)


def marker(sl, x, y, n, d=0.3, color=L700, tc=WHITE, size=10):
    return oval(sl, x, y, d, color, paras=[(str(n), dict(size=size, bold=True, color=tc, align=CEN))], pad=0,
                anchor=MIDDLE)


def tick(sl, x, y, s=0.22, color=L700, lw=2.0):
    line(sl, x, y + s * 0.55, x + s * 0.38, y + s * 0.9, color, lw)
    line(sl, x + s * 0.38, y + s * 0.9, x + s, y + s * 0.1, color, lw)


def cross(sl, x, y, s=0.2, color=RED, lw=2.0):
    line(sl, x, y, x + s, y + s, color, lw)
    line(sl, x + s, y, x, y + s, color, lw)


# -------------------------------------------------------------- composites ---
def caption(sl, x, y, w, text):
    """Exhibit caption above a figure."""
    EXH[0] += 1
    return tb(sl, x, y, w, 0.28, [[('Exhibit %d   ' % EXH[0], dict(size=9, bold=True, color=L700)),
                                   (text, dict(size=9.5, bold=True, color=INK))]])


def heading(sl, x, y, w, text, size=13):
    return tb(sl, x, y, w, 0.32, text, size=size, color=L700, font=DISPLAY)


def notes(sl, x, y, w, h, items, size=10.5, gap=6, head_color=L700):
    """Commentary column: (head, body) pairs as run-in paragraphs, or plain strings."""
    paras = []
    for it in items:
        if isinstance(it, tuple):
            hd, bd = it
            paras.append(([(hd + '  ', dict(size=size, bold=True, color=head_color)), (bd, dict(size=size))],
                          dict(gap=gap, spacing=1.07)))
        else:
            paras.append((it, dict(size=size, gap=gap, spacing=1.07)))
    return tb(sl, x, y, w, h, paras)


def bullets(sl, x, y, w, h, items, size=10.5, gap=5, **kw):
    return tb(sl, x, y, w, h, items, size=size, bullet=True, gap=gap, spacing=1.06, **kw)


def ledger(sl, x, y, w, rows, widths, size=10.5, rh=0.36, head=None, bold_first=True, first_color=L700,
           aligns=None, top_rule=True, colors=None, hsize=9.5, row_heights=None, pad_x=0.06, icons=None,
           icon_color=L700):
    """Rows of text separated by hairlines, drawn with text boxes: a ledger, not a table object."""
    if icons:
        isz = min(0.26, rh * 0.62)
        yy0 = y + (0.32 if head else 0)
        for i, ic in enumerate(icons):
            hh = row_heights[i] if row_heights else rh
            yi = yy0 + (sum(row_heights[:i]) if row_heights else i * rh)
            if ic:
                ic_col = icon_color
                if isinstance(ic, tuple):
                    ic, ic_col = ic
                icon(sl, ic, x, yi + (hh - isz) / 2, isz, ic_col)
        x, w = x + isz + 0.12, w - isz - 0.12
    tot = float(sum(widths))
    xs = [x]
    for cw_ in widths[:-1]:
        xs.append(xs[-1] + w * cw_ / tot)
    ws = [w * cw_ / tot for cw_ in widths]
    yy = y
    if head:
        for j, t in enumerate(head):
            tb(sl, xs[j] + pad_x, yy, ws[j] - 2 * pad_x, 0.3, t, size=hsize, bold=True, color=GREY,
               align=(aligns[j] if aligns else L), anchor=BOTTOM)
        yy += 0.32
        rule(sl, x, yy, w, L700, 1.0)
    elif top_rule:
        rule(sl, x, yy, w, L700, 1.0)
    for i, row in enumerate(rows):
        hh = row_heights[i] if row_heights else rh
        for j, t in enumerate(row):
            col = (colors[i][j] if colors and colors[i] and colors[i][j] else
                   (first_color if (bold_first and j == 0) else INK))
            tb(sl, xs[j] + pad_x, yy + 0.05, ws[j] - 2 * pad_x, hh - 0.08, t, size=size,
               bold=bold_first and j == 0, color=col, align=(aligns[j] if aligns else L), anchor=MIDDLE,
               spacing=1.03)
        yy += hh
        rule(sl, x, yy, w, RULE, 0.75)
    return yy


def table(sl, x, y, w, rows, widths, size=10, hsize=9.5, rh=0.32, aligns=None, bold_first=True, shade=None):
    """Real table object with hairline rules (No Style, No Grid)."""
    nr, nc = len(rows), len(rows[0])
    g = sl.shapes.add_table(nr, nc, E(x), E(y), E(w), E(rh * nr))
    t = g.table
    t.first_row = False
    t.horz_banding = False
    sid = t._tbl.tblPr.find(qn('a:tableStyleId'))
    if sid is None:
        sid = etree.SubElement(t._tbl.tblPr, qn('a:tableStyleId'))
    sid.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
    tot = float(sum(widths))
    for j, cw_ in enumerate(widths):
        t.columns[j].width = E(w * cw_ / tot)
    for i in range(nr):
        t.rows[i].height = E(rh)
        for j in range(nc):
            c = t.cell(i, j)
            c.vertical_anchor = MIDDLE
            c.margin_left = c.margin_right = E(0.06)
            c.margin_top = c.margin_bottom = E(0.03)
            hd = i == 0
            tcPr = c._tc.get_or_add_tcPr()
            if shade and shade.get(i):
                sf0 = etree.SubElement(tcPr, qn('a:solidFill'))
                etree.SubElement(sf0, qn('a:srgbClr')).set('val', str(shade[i]))
            lnB = etree.Element(qn('a:lnB'))
            lnB.set('w', str(int((1.0 if hd else 0.6) * 12700)))
            sf = etree.SubElement(lnB, qn('a:solidFill'))
            etree.SubElement(sf, qn('a:srgbClr')).set('val', str(L700 if hd else RULE))
            tcPr.insert(0, lnB)
            fill(c.text_frame, [rows[i][j]], size=hsize if hd else size, bold=hd or (bold_first and j == 0),
                 color=GREY if hd else (L700 if (bold_first and j == 0) else INK),
                 align=(aligns[j] if aligns else L), spacing=1.0)
    return g


def flow(sl, x, y, w, h, items, gap=0.3, size=10, head_size=11.5, fills=None, lines=None):
    """Boxes in a row joined by arrows. items: (head, body)."""
    n = len(items)
    bw = (w - (n - 1) * gap) / n
    for i, (hd, bd) in enumerate(items):
        xx = x + i * (bw + gap)
        col = fills[i] if fills else WHITE
        ln = lines[i] if lines else RULE
        box(sl, xx, y, bw, h, [(hd, dict(size=head_size, bold=True, color=L700 if col != LIME else INK, gap=3)),
                               (bd, dict(size=size, spacing=1.05))], color=col, line=ln, pad=0.1)
        if i < n - 1:
            line(sl, xx + bw + 0.03, y + h / 2, xx + bw + gap - 0.03, y + h / 2, L700, 1.25, arrow=True)
    return bw


def chart(sl, kind, x, y, w, h, cats, series, colors=None, point_colors=None, fmt='#,##0', labels=True,
          legend=False, size=10, gap=55, overlap=None, val_axis=False, vmax=None, vmin=None, label_pos=None,
          reverse=None, smooth=False, line_w=2.25, grid=False, val_fmt=None, cat_labels=True, cat_low=False):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    gf = sl.shapes.add_chart(kind, E(x), E(y), E(w), E(h), cd)
    ch = gf.chart
    ch.has_title = False
    ch.font.size = Pt(size)
    ch.font.name = BODY
    ch.font.color.rgb = INK
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size - 0.5)
    plot = ch.plots[0]
    is_line = kind in (XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS)
    if not is_line:
        plot.gap_width = gap
        if overlap is not None:
            plot.overlap = overlap
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = fmt
        dl.number_format_is_linked = False
        dl.font.size = Pt(size - 0.5)
        dl.font.color.rgb = INK
        if label_pos is not None:
            dl.position = label_pos
    for i, s in enumerate(plot.series):
        col = colors[i] if colors else L700
        if is_line:
            s.format.line.color.rgb = col
            s.format.line.width = Pt(line_w)
            s.smooth = smooth
        else:
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = col
            s.format.line.fill.background()
            s.invert_if_negative = False
        if point_colors and i == 0:
            for j, pc in enumerate(point_colors):
                pt = s.points[j]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = pc
            for dPt in s._element.findall(qn('c:dPt')):
                inv = dPt.find(qn('c:invertIfNegative'))
                if inv is None:
                    inv = etree.Element(qn('c:invertIfNegative'))
                    dPt.insert(1, inv)
                inv.set('val', '0')
    va = ch.value_axis
    va.has_major_gridlines = grid
    if grid:
        va.major_gridlines.format.line.color.rgb = C('E4E9DF')
        va.major_gridlines.format.line.width = Pt(0.5)
    va.visible = val_axis
    if val_axis:
        va.tick_labels.font.size = Pt(size - 1)
        va.tick_labels.font.color.rgb = GREY
        va.format.line.fill.background()
        va.major_tick_mark = XL_TICK_MARK.NONE
        if val_fmt:
            va.tick_labels.number_format = val_fmt
            va.tick_labels.number_format_is_linked = False
    if vmax is not None:
        va.maximum_scale = vmax
    if vmin is not None:
        va.minimum_scale = vmin
    ca = ch.category_axis
    ca.format.line.color.rgb = GREY_LT
    ca.has_major_gridlines = False
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.tick_labels.font.size = Pt(size)
    if not cat_labels:
        ca.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    elif is_line or cat_low:
        ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    if reverse is None:
        reverse = kind in (XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.BAR_STACKED)
    if reverse:
        ca.reverse_order = True
    return ch


def pic(sl, path, x, y, h=None, w=None):
    kw = {}
    if h is not None:
        kw['height'] = E(h)
    if w is not None:
        kw['width'] = E(w)
    return sl.shapes.add_picture(path, E(x), E(y), **kw)


# ------------------------------------------------------------------ chrome ---
def slide(title, lead=None, src=None, partner=False, lead_size=12.5):
    sl = prs.slides.add_slide(BLANK)
    PAGE[0] += 1
    TOC.append((PAGE[0], title, SECTION[0]))
    tb(sl, M, 0.36, 10.3, 0.5, title, size=24, color=L700, font=DISPLAY)
    if lead:
        tb(sl, M, 0.9, 10.5, 0.66, lead, size=lead_size, color=INK, spacing=1.06)
    lg = pic(sl, LOGO, 0, 0.44, h=0.3)
    lg.left = E(W - M) - lg.width
    if partner:
        mq = pic(sl, MQ_LOGO, 0, 0.3, h=0.58)
        mq.left = lg.left - E(0.28) - mq.width
        vrule(sl, (lg.left - E(0.14)) / 914400.0, 0.36, 0.46, GREY_LT, 0.75)
    tb(sl, M, 7.08, 8, 0.22, 'Halqa   Complete Position   ' + DATE + ('   ' + SECTION[0] if SECTION[0] else ''),
       size=8, color=GREY)
    tb(sl, W - M - 1.0, 7.08, 1.0, 0.22, str(PAGE[0]), size=8, color=GREY, align=R)
    if src:
        tb(sl, M, YS, CW, 0.3, 'Source: ' + src if not src.startswith(('Source', 'Sources', 'Note', 'Modelled',
                                                                          'Assumptions')) else src,
           size=8, color=GREY, spacing=1.0)
    return sl


DIVIDERS = []


def divider(num, title, blurb):
    SECTION[0] = title
    sl = prs.slides.add_slide(BLANK)
    PAGE[0] += 1
    rect(sl, 0, 0, 4.7, H, LIME)
    tb(sl, 0.62, 0.9, 3.6, 1.6, num, size=96, color=WHITE, font=DISPLAY)
    tb(sl, 0.66, 2.75, 3.7, 1.6, title, size=30, color=INK, font=DISPLAY, spacing=0.95)
    tb(sl, 0.68, 4.45, 3.55, 2.2, blurb, size=12, color=INK, spacing=1.12)
    lg = pic(sl, LOGO, 0, 0.44, h=0.3)
    lg.left = E(W - M) - lg.width
    tb(sl, 5.3, 1.02, 6, 0.3, 'In this section', size=11, bold=True, color=GREY)
    tb(sl, W - M - 1.0, 7.08, 1.0, 0.22, str(PAGE[0]), size=8, color=GREY, align=R)
    DIVIDERS.append((sl, title, PAGE[0]))
    return sl


def finish_dividers():
    """Write the slide list on each divider once every page number is known."""
    for sl, title, pg in DIVIDERS:
        items = [(p, t) for p, t, s in TOC if s == title]
        y = 1.42
        rh = min(0.42, 5.3 / max(1, len(items)))
        rule(sl, 5.3, y, W - M - 5.3, L700, 1.0)
        for p, t in items:
            tb(sl, 5.3, y + 0.02, 6.4, rh - 0.04, t, size=12.5 if rh >= 0.38 else 11.5, color=INK, anchor=MIDDLE)
            tb(sl, W - M - 0.8, y + 0.02, 0.8, rh - 0.04, str(p), size=11, color=GREY, align=R, anchor=MIDDLE)
            y += rh
            rule(sl, 5.3, y, W - M - 5.3, RULE, 0.75)


def check_and_save(path=OUT):
    cp = prs.core_properties
    cp.title = 'Halqa: Complete Position'
    cp.author = 'Halqa'
    cp.last_modified_by = 'Halqa'
    prs.save(path)
    bad = []
    for n, s in enumerate(prs.slides, 1):
        for shp in s.shapes:
            texts = []
            if shp.has_text_frame:
                texts.append(shp.text_frame.text)
            if getattr(shp, 'has_table', False) and shp.has_table:
                texts += [c.text_frame.text for r_ in shp.table.rows for c in r_.cells]
            for t in texts:
                for pat in (r'\byou\b', r'\byour\b', r'\bwe\b', r'\bour\b', u'[‒–—―]', r' - ',
                            r'(?i)\bjourney\b|\bunlock|\bseamless|\bempower|\bleverag|\brevolution|\bgame.?changer'
                            r'|(?<!Business )\brecorder\b|\bSakh\b|\btruly\b'):
                    if re.search(pat, t):
                        bad.append((n, pat, t[:80]))
    print('slides', len(prs.slides), 'bytes', os.path.getsize(path))
    for b in bad:
        print('CHECK', b)


# ------------------------------------------------------------------ icons ---
import json as _json
ICON_SRC = r'D:/HALQA SIGMA APP/halqa-web/node_modules/lucide-react/dist/esm/icons'
ICON_DIR = os.path.join(HERE, '..', 'icons')
os.makedirs(ICON_DIR, exist_ok=True)


def _icon_svg(name, color, sw):
    s = open(os.path.join(ICON_SRC, name + '.mjs'), encoding='utf-8').read()
    js = re.search(r'const __iconNode = (\[.*?\]);\n', s, re.S).group(1)
    js = re.sub(r'(\{|,)\s*([A-Za-z_][\w-]*)\s*:', r'\1 "\2":', js)
    parts = ['<%s %s/>' % (tag, ' '.join('%s="%s"' % (k, v) for k, v in attrs.items() if k != 'key'))
             for tag, attrs in _json.loads(js)]
    return ('<svg xmlns="http://www.w3.org/2000/svg" width="24" height="24" viewBox="0 0 24 24" fill="none" '
            'stroke="#%s" stroke-width="%s" stroke-linecap="round" stroke-linejoin="round">%s</svg>'
            % (color, sw, ''.join(parts)))


def icon_png(name, color='41801A', sw=2.0, px=256):
    path = os.path.join(ICON_DIR, '%s-%s-%s.png' % (name, color, sw))
    if not os.path.exists(path):
        import pymupdf
        doc = pymupdf.open(stream=_icon_svg(name, color, sw).encode(), filetype='svg')
        doc[0].get_pixmap(matrix=pymupdf.Matrix(px / 24.0, px / 24.0), alpha=True).save(path)
    return path


def icon(sl, name, x, y, s=0.36, color=L700, sw=2.0):
    return pic(sl, icon_png(name, str(color), sw), x, y, h=s)


def icon_disc(sl, name, x, y, d=0.5, fill_=LIME_XLT, color=L700, sw=2.0, line_=None):
    """Icon centred in a filled circle."""
    oval(sl, x, y, d, fill_, line=line_)
    s = d * 0.56
    return icon(sl, name, x + (d - s) / 2, y + (d - s) / 2, s, color, sw)


# ---------------------------------------------------------- phone mockups ---
SHOTS = os.path.join(HERE, '..', 'shots2', 'light')


def phone_png(name, crop=None):
    """Screenshot in a plain dark phone frame, cached as PNG."""
    out = os.path.join(SHOTS, 'phone-%s-%s.png' % (name, crop or 'full'))
    if not os.path.exists(out):
        from PIL import Image, ImageDraw
        im = Image.open(os.path.join(SHOTS, name + '.png')).convert('RGB')
        if crop:
            im = im.crop((0, 0, im.width, crop))
        pad, r_out = 34, 150
        W_, H_ = im.width + 2 * pad, im.height + 2 * pad
        frame = Image.new('RGBA', (W_, H_), (0, 0, 0, 0))
        d = ImageDraw.Draw(frame)
        d.rounded_rectangle((0, 0, W_ - 1, H_ - 1), radius=r_out, fill=(27, 31, 24, 255))
        mask = Image.new('L', im.size, 0)
        ImageDraw.Draw(mask).rounded_rectangle((0, 0, im.width - 1, im.height - 1), radius=r_out - pad, fill=255)
        frame.paste(im, (pad, pad), mask)
        frame = frame.resize((frame.width * 7 // 12, frame.height * 7 // 12), Image.LANCZOS)
        frame.save(out, optimize=True)
    return out


def phone(sl, name, x, y, h, crop=None):
    return pic(sl, phone_png(name, crop), x, y, h=h)


def icon_rows(sl, x, y, w, rows, rh=0.62, size=10.5, isz=0.34, head_w=None, rules=True, disc=False,
              head_color=L700):
    """Rows of (icon, head, body): an icon, a bold head and a line of text, with hairlines between rows."""
    for i, (ic, hd, bd) in enumerate(rows):
        yy = y + i * rh
        if disc:
            icon_disc(sl, ic, x, yy + (rh - isz * 1.3) / 2, isz * 1.3)
            tx = x + isz * 1.3 + 0.18
        else:
            icon(sl, ic, x, yy + (rh - isz) / 2, isz)
            tx = x + isz + 0.18
        if head_w:
            tb(sl, tx, yy, head_w, rh, hd, size=size, bold=True, color=head_color, anchor=MIDDLE, spacing=1.03)
            tb(sl, tx + head_w + 0.1, yy, w - (tx - x) - head_w - 0.1, rh, bd, size=size - 0.5, anchor=MIDDLE,
               spacing=1.04)
        else:
            tb(sl, tx, yy, w - (tx - x), rh, [[(hd + '  ', dict(size=size, bold=True, color=head_color)),
                                               (bd, dict(size=size - 0.5))]], anchor=MIDDLE, spacing=1.05)
        if rules and i < len(rows) - 1:
            rule(sl, tx, yy + rh, w - (tx - x))
    return y + len(rows) * rh


def stat_rows(sl, x, y, w, rows, rh=0.6, num_w=1.25, isz=0.34, size=10.5, nsize=17):
    """Rows of (icon, figure, label): an infographic list without boxes."""
    for i, (ic, n_, lab) in enumerate(rows):
        yy = y + i * rh
        icon(sl, ic, x, yy + (rh - isz) / 2, isz)
        tb(sl, x + isz + 0.15, yy, num_w, rh, n_, size=nsize, color=L700, font=DISPLAY, anchor=MIDDLE)
        tb(sl, x + isz + 0.2 + num_w, yy, w - isz - 0.2 - num_w, rh, lab, size=size, anchor=MIDDLE, spacing=1.04)
        if i < len(rows) - 1:
            rule(sl, x, yy + rh, w)
    return y + len(rows) * rh
