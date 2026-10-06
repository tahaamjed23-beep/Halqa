# -*- coding: utf-8 -*-
"""
Halqa: the complete position, one deck for everything (28 September 2026).

Built on the bank route decided on 28 September 2026 (Mashreq Bank Pakistan as partner bank, decisions D1 to D16).
Lime ramp only (l500 6DC72A fills, l700 41801A text), no pine. Diagrams in place of most tables. Plain wording.
Third person, no dashes, noun headings. Sources in the footer notes.
"""
import io, math, os, re, zipfile
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, 'lockup-520.png')
MARK = os.path.join(HERE, 'mark-160.png')
OUT = os.path.join(HERE, 'Halqa-Complete-Position.pptx')
DATE = '28 September 2026'


def C(h):
    return RGBColor.from_string(h)


LIME, L700, L600 = C('6DC72A'), C('41801A'), C('58A822')
LIME_LT, LIME_XLT = C('E3F8CA'), C('F3FCE7')
INK, GREY, GREY_LT, RULE, WHITE = C('0C1408'), C('5F6B58'), C('A9B4A3'), C('C9D6BF'), C('FFFFFF')
RED, AMBER = C('C0392B'), C('D9A21B')
DISPLAY, BODY = 'Georgia', 'Arial'
L, R, CEN = PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.CENTER
TOP, MIDDLE = MSO_ANCHOR.TOP, MSO_ANCHOR.MIDDLE
W, H, M = 13.333, 7.5, 0.62
CW = W - 2 * M


def E(v):
    return int(round(v * 914400))


prs = Presentation()
prs.slide_width, prs.slide_height = E(W), E(H)
BLANK = [lay for lay in prs.slide_layouts if lay.name == 'Blank'][0]
for lay in list(prs.slide_layouts):
    if lay.name != 'Blank':
        prs.slide_layouts.remove(lay)
PAGE = [0]
SECTION = ['']


# ------------------------------------------------------------------- text ---
def _runs(p, text, st):
    segs = []
    if isinstance(text, list):
        segs = [(t, dict(st, **s2)) for t, s2 in text]
    else:
        for k, seg in enumerate(text.split('**')):
            if seg:
                segs.append((seg, dict(st, bold=st.get('bold') or k % 2 == 1)))
    for seg, s in segs:
        r = p.add_run()
        r.text = seg.upper() if s.get('caps') else seg
        f = r.font
        f.size = Pt(s.get('size', 12))
        f.name = s.get('font', BODY)
        f.color.rgb = s.get('color', INK)
        if s.get('bold'):
            f.bold = True
        if s.get('italic'):
            f.italic = True
        if s.get('track'):
            r.font._rPr.set('spc', str(int(s['track'] * 100)))


def fill(tf, paras, **st):
    tf.word_wrap = True
    if isinstance(paras, str):
        paras = [paras]
    for i, item in enumerate(paras):
        text, own = (item if isinstance(item, tuple) else (item, {}))
        s = dict(st, **own)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if s.get('align', L) != L:
            p.alignment = s['align']
        elif p._p.pPr is not None and p._p.pPr.get('algn') is not None:
            del p._p.pPr.attrib['algn']
        if s.get('spacing'):
            p.line_spacing = s['spacing']
        if s.get('gap'):
            p.space_after = Pt(s['gap'])
        if s.get('bullet') or s.get('num'):
            pPr = p._p.get_or_add_pPr()
            ind = E(s.get('ind', 0.24))
            pPr.set('marL', str(ind))
            pPr.set('indent', str(-ind))
            bc = etree.SubElement(pPr, qn('a:buClr'))
            etree.SubElement(bc, qn('a:srgbClr')).set('val', str(s.get('bcolor', LIME)))
            if s.get('bullet'):
                etree.SubElement(pPr, qn('a:buSzPct')).set('val', '75000')
                etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
                etree.SubElement(pPr, qn('a:buChar')).set('char', u'■')
            else:
                etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
                etree.SubElement(pPr, qn('a:buAutoNum')).set('type', 'arabicPeriod')
        _runs(p, text, s)


def _frame(tf, pad, anchor, force=False):
    if isinstance(pad, (int, float)):
        pad = (pad, pad * 0.75, pad, pad * 0.75)
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [E(v) for v in pad]
    if anchor != TOP or force:
        tf.vertical_anchor = anchor


def tb(sl, x, y, w, h, paras, pad=0, anchor=TOP, **st):
    box = sl.shapes.add_textbox(E(x), E(y), E(w), E(h))
    _frame(box.text_frame, pad, anchor)
    fill(box.text_frame, paras, **st)
    return box


def _nostyle(shape):
    st = shape._element.find(qn('p:style'))
    if st is not None:
        shape._element.remove(st)


def shape(sl, kind, x, y, w, h, color=None, line=None, lw=0.75, rot=0, paras=None, pad=0.16, anchor=TOP, **st):
    s = sl.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
    _nostyle(s)
    if color is not None:
        s.fill.solid()
        s.fill.fore_color.rgb = color
    else:
        s.fill.background()
    if line is not None:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    if rot:
        s.rotation = rot
    if paras is not None:
        _frame(s.text_frame, pad, anchor, force=True)
        fill(s.text_frame, paras, **st)
    return s


def rect(sl, x, y, w, h, color=None, **kw):
    return shape(sl, MSO_SHAPE.RECTANGLE, x, y, w, h, color, **kw)


def line(sl, x1, y1, x2, y2, color=L700, lw=1.0, arrow=False, dash=False):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    _nostyle(c)
    c.line.color.rgb = color
    c.line.width = Pt(lw)
    ln = c.line._get_or_add_ln()
    if dash:
        etree.SubElement(ln, qn('a:prstDash')).set('val', 'dash')
    if arrow:
        te = etree.SubElement(ln, qn('a:tailEnd'))
        te.set('type', 'triangle')
        te.set('w', 'med')
        te.set('len', 'med')
    return c


def table(sl, x, y, w, rows, widths, size=11, hsize=10.5, rh=0.36, aligns=None, bold_first=True):
    nr, nc = len(rows), len(rows[0])
    g = sl.shapes.add_table(nr, nc, E(x), E(y), E(w), E(rh * nr))
    t = g.table
    t.first_row = False
    t.horz_banding = False
    sid = t._tbl.tblPr.find(qn('a:tableStyleId'))
    if sid is None:
        sid = etree.SubElement(t._tbl.tblPr, qn('a:tableStyleId'))
    sid.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
    tot = float(sum(widths))
    for j, cw in enumerate(widths):
        t.columns[j].width = E(w * cw / tot)
    for i in range(nr):
        t.rows[i].height = E(rh)
        for j in range(nc):
            c = t.cell(i, j)
            c.vertical_anchor = MIDDLE
            head = i == 0
            tcPr = c._tc.get_or_add_tcPr()
            lnB = etree.SubElement(tcPr, qn('a:lnB'))
            lnB.set('w', str(int((1.25 if head else 0.75) * 12700)))
            sf = etree.SubElement(lnB, qn('a:solidFill'))
            etree.SubElement(sf, qn('a:srgbClr')).set('val', str(L700 if head else RULE))
            fill(c.text_frame, [rows[i][j]], size=hsize if head else size, bold=head or (bold_first and j == 0),
                 color=L700 if head else INK, align=(aligns[j] if aligns else L))
    return g


# ------------------------------------------------------------------ marks ---
def tick(sl, x, y, s=0.26, color=L700, lw=2.4):
    line(sl, x, y + s * 0.55, x + s * 0.38, y + s * 0.9, color, lw)
    line(sl, x + s * 0.38, y + s * 0.9, x + s, y + s * 0.1, color, lw)


def cross(sl, x, y, s=0.26, color=RED, lw=2.4):
    line(sl, x, y, x + s, y + s, color, lw)
    line(sl, x + s, y, x, y + s, color, lw)


def icon_bank(sl, x, y, s=0.6, color=L700):
    shape(sl, MSO_SHAPE.ISOSCELES_TRIANGLE, x, y, s, s * 0.32, color)
    for k in range(3):
        rect(sl, x + s * (0.12 + k * 0.3), y + s * 0.38, s * 0.16, s * 0.42, color)
    rect(sl, x, y + s * 0.84, s, s * 0.14, color)


def icon_phone(sl, x, y, s=0.6, color=L700):
    rect(sl, x + s * 0.22, y, s * 0.56, s, None, line=color, lw=2.0)
    rect(sl, x + s * 0.42, y + s * 0.82, s * 0.16, s * 0.08, color)


def icon_person(sl, x, y, s=0.6, color=L700):
    shape(sl, MSO_SHAPE.OVAL, x + s * 0.3, y, s * 0.4, s * 0.4, color)
    rect(sl, x + s * 0.12, y + s * 0.48, s * 0.76, s * 0.5, color)


def icon_doc(sl, x, y, s=0.6, color=L700):
    rect(sl, x + s * 0.15, y, s * 0.7, s, None, line=color, lw=2.0)
    for k in range(3):
        rect(sl, x + s * 0.28, y + s * (0.22 + k * 0.22), s * 0.44, s * 0.06, color)


def icon_group(sl, x, y, s=0.6, color=L700):
    for k in range(3):
        icon_person(sl, x + k * s * 0.34, y + (0.0 if k == 1 else s * 0.12), s * 0.42, color)


def ring(sl, cx, cy, r, bw, bh, lit=0, on=INK, off=WHITE, n=12):
    for k in range(n):
        a = math.radians(k * 360.0 / n)
        px, py = cx + r * math.sin(a), cy - r * math.cos(a)
        rect(sl, px - bw / 2, py - bh / 2, bw, bh, on if k == lit else off, rot=k * 360.0 / n)


def hbars(sl, x, y, w, items, bar_h=0.34, gap=0.56, label_w=1.9, val_w=1.1, maxv=None, size=11.5):
    maxv = maxv or max(i[1] for i in items)
    px = x + label_w
    pw = w - label_w - val_w
    for i, (lab, v, disp, col) in enumerate(items):
        yy = y + i * gap
        tb(sl, x, yy + 0.02, label_w - 0.12, bar_h, lab, size=size, align=R, anchor=MIDDLE)
        bw = max(0.03, pw * float(v) / maxv)
        rect(sl, px, yy, bw, bar_h, col)
        tb(sl, px + bw + 0.08, yy + 0.02, val_w, bar_h, disp, size=size, color=L700, bold=True, anchor=MIDDLE)
    line(sl, px, y - 0.06, px, y + gap * (len(items) - 1) + bar_h + 0.06, INK, 0.75)


def flow(sl, y, h, items, x0=M, width=CW, gap=0.34, fill_=LIME_XLT, size=11, title_size=13.5):
    n = len(items)
    bw = (width - (n - 1) * gap) / n
    for i, it in enumerate(items):
        x = x0 + i * (bw + gap)
        head, body = it[0], it[1]
        col = it[2] if len(it) > 2 else fill_
        rect(sl, x, y, bw, h, col, pad=0.14, paras=[
            (head, dict(size=title_size, font=DISPLAY, color=L700 if col != LIME else INK, gap=4)),
            (body, dict(size=size, spacing=1.08))])
        if i < n - 1:
            line(sl, x + bw + 0.04, y + 0.36, x + bw + gap - 0.04, y + 0.36, L700, 1.5, arrow=True)
    return bw


# ------------------------------------------------------------------ chrome ---
def slide(title, lead=None, lead_size=14):
    sl = prs.slides.add_slide(BLANK)
    PAGE[0] += 1
    tb(sl, M, 0.5, 10.5, 0.62, title, size=28, color=L700, font=DISPLAY)
    if lead:
        tb(sl, M, 1.12, 11.0, 0.62, lead, size=lead_size, color=INK, spacing=1.08)
    pic = sl.shapes.add_picture(LOGO, 0, E(0.5), height=E(0.34))
    pic.left = E(W - M) - pic.width
    tb(sl, M, 7.08, 9, 0.22, 'Halqa   ' + SECTION[0], size=8, color=GREY)
    tb(sl, W - M - 1.0, 7.08, 1.0, 0.22, str(PAGE[0]), size=8, color=GREY, align=R)
    return sl


def note(sl, text, y=6.66):
    tb(sl, M, y, CW, 0.36, text, size=8.5, color=GREY, spacing=1.05)


def divider(num, title, items):
    SECTION[0] = title
    sl = prs.slides.add_slide(BLANK)
    PAGE[0] += 1
    rect(sl, 0, 0, W, H, LIME)
    tb(sl, M + 0.1, 1.2, 3, 1.6, num, size=110, color=WHITE, font=DISPLAY)
    tb(sl, M + 0.15, 3.0, 8.5, 0.9, title, size=40, color=INK, font=DISPLAY)
    tb(sl, M + 0.18, 4.05, 8.0, 2.5, items, size=14, color=INK, gap=6, spacing=1.1)
    ring(sl, 10.35, 3.75, 1.5, 0.2, 0.6, lit=0, on=INK, off=WHITE)
    tb(sl, W - M - 1.0, 7.08, 1.0, 0.22, str(PAGE[0]), size=8, color=INK, align=R)
    return sl


# =================================================================== COVER ===
sl = prs.slides.add_slide(BLANK)
PAGE[0] += 1
rect(sl, 8.3, 0, W - 8.3, H, LIME)
ring(sl, 8.3 + (W - 8.3) / 2, 3.6, 1.45, 0.2, 0.6)
sl.shapes.add_picture(LOGO, E(0.9), E(0.8), height=E(0.6))
tb(sl, 0.9, 2.45, 7.2, 1.3, ['The complete', 'position'], size=44, color=INK, font=DISPLAY, spacing=0.92)
tb(sl, 0.9, 4.05, 6.9, 1.2, 'Committee savings run by Halqa, with the money held and moved by a partner bank. '
   'Everything in one deck: the committee, the partnership with Mashreq Bank Pakistan, the product, the '
   'evidence, the law and the plan.', size=15, color=GREY, spacing=1.15)
tb(sl, 0.9, 6.05, 6, 0.3, 'Taha Amjed, Chairman', size=12, color=INK, bold=True)
tb(sl, 0.9, 6.38, 6, 0.3, DATE, size=10, color=GREY)
tb(sl, 8.3, 5.6, W - 8.3, 0.3, 'Twelve members, one pot a month', size=10.5, color=INK, align=CEN)

# ================================================================ CONTENTS ===
SECTION[0] = 'Contents'
sl = slide('Contents')
parts = [('1', 'The committee', 'What it is, where the risk sits, who uses it', '3'),
         ('2', 'The bank partnership', 'Why a bank, roles, Mashreq, the money route, fees', '9'),
         ('3', 'The product', 'Signup to completion, bands, cover, default, credit, points, Hyper', '18'),
         ('4', 'Evidence and competition', 'Twenty five attempts, the models that work, Oraan, JazzCash', '29'),
         ('5', 'Law, risk and plan', 'Legal map, controls, integration, approvals, pilot, open questions', '35')]
y = 1.55
for n, t, d, pg in parts:
    rect(sl, M, y, 0.72, 0.72, LIME, paras=[(n, dict(size=24, font=DISPLAY, color=INK, align=CEN))],
         pad=0, anchor=MIDDLE)
    tb(sl, M + 1.0, y + 0.02, 5.2, 0.4, t, size=20, color=L700, font=DISPLAY)
    tb(sl, M + 1.0, y + 0.42, 8.5, 0.3, d, size=12, color=GREY)
    tb(sl, W - M - 1.0, y + 0.1, 1.0, 0.4, 'page ' + pg, size=11, color=GREY, align=R)
    y += 1.02

# ============================================================ 1 COMMITTEE ===
divider('1', 'The committee', ['What a committee is', 'Where its risk sits', 'How the order is set',
                               'The market, and who uses it'])

# ---- what it is
sl = slide('What a Committee Is',
           'Each member pays the same amount every month and each month one member takes the whole pot. By the end, '
           'every member has paid in exactly what that member took out.')
cx, cy, r = 3.25, 4.25, 1.75
for k in range(12):
    a = math.radians(k * 30)
    px, py = cx + r * math.sin(a), cy - r * math.cos(a)
    rect(sl, px - 0.33, py - 0.24, 0.66, 0.48, LIME if k == 0 else LIME_XLT, line=L700 if k == 0 else RULE,
         paras=[('Seat %d' % (k + 1), dict(size=9.5, bold=k == 0, color=INK, align=CEN))], pad=0, anchor=MIDDLE)
tb(sl, cx - 1.1, cy - 0.45, 2.2, 0.9, [('Pot', dict(size=12, color=GREY, align=CEN)),
                                        ('Rs 120,000', dict(size=22, font=DISPLAY, color=L700, align=CEN))])
tb(sl, cx - 1.6, cy + r + 0.42, 3.2, 0.3, 'Month 1: seat 1 collects', size=10.5, color=GREY, align=CEN)
x0 = 6.55
tb(sl, x0, 2.1, 6.2, 0.4, 'Twelve members at Rs 10,000 a month', size=15, color=L700, font=DISPLAY)
rows = [('Seat 1', 'receives Rs 120,000 after paying Rs 10,000, then pays Rs 110,000 more'),
        ('Seat 6', 'receives Rs 120,000 after paying Rs 60,000'),
        ('Seat 12', 'receives Rs 120,000 after paying all Rs 120,000'),
        ('Every seat', 'pays Rs 120,000 and receives Rs 120,000')]
yy = 2.65
for a_, b_ in rows:
    tb(sl, x0, yy, 1.35, 0.5, a_, size=13, color=INK, bold=True)
    tb(sl, x0 + 1.4, yy, 4.8, 0.6, b_, size=13, spacing=1.08)
    line(sl, x0, yy + 0.62, W - M, yy + 0.62, RULE, 0.75)
    yy += 0.74
rect(sl, x0, 5.7, 2.95, 0.95, LIME_XLT, paras=[
    ('Early seats', dict(size=12.5, bold=True, color=L700, gap=2)),
    ('An advance without interest, repaid in instalments', dict(size=11))])
rect(sl, x0 + 3.1, 5.7, 3.1, 0.95, LIME_XLT, paras=[
    ('Late seats', dict(size=12.5, bold=True, color=L700, gap=2)),
    ('Saving under a commitment the group enforces', dict(size=11))])

# ---- risk
sl = slide('Where the Risk Sits',
           'A member who collects early still owes the rest of the instalments. That amount is largest at seat 1 and '
           'nil at seat 12, so the early seats are where a default can hurt the circle.')
items = []
for s_ in range(1, 13):
    v = 10000 * (12 - s_)
    col = L700 if s_ <= 6 else (LIME if s_ <= 9 else LIME_LT)
    items.append(('Seat %d' % s_, v, 'Rs {:,}'.format(v) if v else 'nil', col))
tb(sl, M, 1.95, 6.6, 0.3, 'Still owed after collecting, twelve members at Rs 10,000', size=11, color=GREY)
hbars(sl, M, 2.35, 6.6, items, bar_h=0.26, gap=0.35, label_w=0.95, val_w=1.15, maxv=110000, size=10.5)
x0 = 7.75
tb(sl, x0, 1.95, 4.95, 0.4, 'Who may take an early seat', size=16, color=L700, font=DISPLAY)
tb(sl, x0, 2.45, 4.95, 3.9, [
    '**Credit bands.** Seats 1 to 6 need the Good or Excellent band; seats 7 to 9 are open from Fair; seats 10 to '
    '12 are open to every band.',
    '**New members.** Only the last three seats, until two circles are completed cleanly and a verification call '
    'is made.',
    '**Affordability.** All committee instalments together stay within a third of verified income.',
    '**Cover.** A default after collecting is covered by the option Mashreq chooses: its own guarantee, insurance '
    'or takaful.',
    '**Halqa takes no seat** and advances no money.'], size=12, bullet=True, gap=9, spacing=1.08)
rect(sl, M + 0.95, 6.62, 0.22, 0.16, L700)
tb(sl, M + 1.22, 6.58, 1.9, 0.25, 'Good and Excellent', size=9, color=GREY)
rect(sl, M + 3.0, 6.62, 0.22, 0.16, LIME)
tb(sl, M + 3.27, 6.58, 1.4, 0.25, 'Fair and above', size=9, color=GREY)
rect(sl, M + 4.6, 6.62, 0.22, 0.16, LIME_LT)
tb(sl, M + 4.87, 6.58, 1.9, 0.25, 'Every band', size=9, color=GREY)

# ---- order
sl = slide('Setting the Order',
           'The order of collection is chosen at joining, within what each member’s band allows, or set by the host. '
           'Two other methods are refused because they change what the arrangement is in law.')
tiles = [('Choice at joining', True, 'The member picks a free seat from those the band allows. The standard method.'),
         ('Set by the host', True, 'The host fixes the order when the circle is created, where the circle allows it.'),
         ('Auction', False, 'The member who accepts the biggest discount collects first. The discount is a price for '
                            'money over time, which is interest. India regulates this separately as chit funds.'),
         ('Prize or lucky draw', False, 'A drawn winner takes the pot and stops paying. Once contributions no longer '
                                        'equal payouts it is a lottery under section 294A of the Penal Code.')]
tw = (CW - 3 * 0.28) / 4
for i, (t, ok, d) in enumerate(tiles):
    x = M + i * (tw + 0.28)
    rect(sl, x, 2.05, tw, 3.7, LIME_XLT if ok else C('F7F7F5'))
    (tick if ok else cross)(sl, x + 0.22, 2.3, 0.34)
    tb(sl, x + 0.7, 2.28, tw - 0.85, 0.4, 'Offered' if ok else 'Refused', size=11, bold=True,
       color=L700 if ok else RED)
    tb(sl, x + 0.22, 2.85, tw - 0.4, 0.8, t, size=17, color=INK, font=DISPLAY)
    tb(sl, x + 0.22, 3.55, tw - 0.4, 2.1, d, size=11.5, spacing=1.12)
rect(sl, M, 6.0, CW, 0.62, LIME_LT, pad=0.2, anchor=MIDDLE, paras=[
    ('Once a circle has started, members may still swap seats with each other in the turn market (page 25), inside '
     'the bands and with Mashreq’s approval of the product.', dict(size=11.5))])

# ---- market
sl = slide('The Market',
           'Committees are the savings habit that already exists outside the banks. They compete with cash kept at '
           'home; the bank account comes a distant third.')
stats = [('Rs 4 trillion', 'rotates through committees each year'), ('52 million', 'adults take part'),
         ('37%', 'of adults have used one'), ('12%', 'have lost money to a committee fraud'),
         ('31%', 'can send or receive a text message')]
sw = CW / 5
for i, (n_, t_) in enumerate(stats):
    x = M + i * sw
    tb(sl, x, 1.95, sw - 0.15, 0.6, n_, size=26, color=L700, font=DISPLAY)
    tb(sl, x, 2.6, sw - 0.25, 0.6, t_, size=11, color=INK, spacing=1.05)
    if i:
        line(sl, x - 0.1, 2.0, x - 0.1, 3.1, RULE, 0.75)
tb(sl, M, 3.55, 6.0, 0.35, 'Where a Pakistani saver keeps money', size=14, color=L700, font=DISPLAY)
hbars(sl, M, 4.05, 6.0, [('Cash at home', 63, '63%', GREY_LT), ('A committee', 33, '33%', LIME),
                         ('A bank account', 13, '13%', L700)], bar_h=0.4, gap=0.62, label_w=1.6, val_w=0.8,
      maxv=70, size=12)
x0 = 7.3
tb(sl, x0, 3.55, 5.4, 0.35, 'The national targets', size=14, color=L700, font=DISPLAY)
tb(sl, x0, 4.05, 5.4, 2.4, [
    'Adults with a financial account: **64 per cent** in 2023, target **75 per cent** by 2028.',
    'Gender gap in financial inclusion: **34 per cent** in 2023, target **25 per cent** by 2028.',
    'Women take part in committees at **twice** the rate of men.',
    'About **100 million** people lack formal financial services.'], size=12, bullet=True, gap=8, spacing=1.08)
note(sl, 'Financial Inclusion Insights; World Bank; National Financial Inclusion Strategy 2024 to 2028 (Profit, 13 '
         'January 2025); Oraan’s figures in Dawn, 27 September 2021. Estimates of participation range from 34 to 41 '
         'per cent across sources.')

# ---- who
sl = slide('Who Uses Committees',
           'Committees select on how regular income is, not on how high it is. Groups exclude people whose income is '
           'irregular, whatever its level.')
gx, gy, gw, gh = M + 0.6, 2.05, 5.6, 4.1
rect(sl, gx, gy, gw, gh / 2, LIME_LT)
rect(sl, gx, gy + gh / 2, gw, gh / 2, C('F7F7F5'))
line(sl, gx, gy + gh, gx + gw + 0.2, gy + gh, INK, 1.0, arrow=True)
line(sl, gx, gy + gh, gx, gy - 0.2, INK, 1.0, arrow=True)
tb(sl, gx + gw - 2.0, gy + gh + 0.08, 2.2, 0.3, 'Income level', size=10.5, color=GREY, align=R)
tb(sl, M - 0.1, gy - 0.05, 0.6, 1.6, 'Regular', size=10.5, color=GREY)
tb(sl, M - 0.1, gy + gh - 0.35, 0.6, 0.3, 'Irregular', size=10.5, color=GREY)
tb(sl, gx + 0.25, gy + 0.2, gw - 0.5, 1.6, [
    ('Served well', dict(size=15, font=DISPLAY, color=L700, gap=4)),
    ('A tailor on Rs 17,000 or a driver on Rs 22,000 a month. Payment is regular and default is rare.',
     dict(size=11.5, spacing=1.1))])
tb(sl, gx + 0.25, gy + gh / 2 + 0.2, gw - 0.5, 1.6, [
    ('Left out by the groups', dict(size=15, font=DISPLAY, color=GREY, gap=4)),
    ('Irregular or insufficient income. Members treat these applicants as bad risks.', dict(size=11.5, spacing=1.1))])
x0 = 7.3
tb(sl, x0, 2.0, 5.4, 4.4, [
    ('Khan, 2013', dict(size=14, font=DISPLAY, color=L700, gap=2)),
    ('Fieldwork in Dera Ghazi Khan. Regular income is a condition of membership, so committees cannot replace '
     'microfinance for the poorest.', dict(size=11.5, spacing=1.1, gap=12)),
    ('Kamran, 2017', dict(size=14, font=DISPLAY, color=L700, gap=2)),
    ('Thirty unbanked informants. None had seen a member default. A housemaid on Rs 15,000 a month called the '
     'instalment the one payment that is never missed, ranked above rent.', dict(size=11.5, spacing=1.1, gap=12)),
    ('What Halqa claims', dict(size=14, font=DISPLAY, color=L700, gap=2)),
    ('A route into formal finance for regular earners without a bank account. Halqa does not claim to serve the '
     'poorest.', dict(size=11.5, spacing=1.1))])

# ==================================================== 2 BANK PARTNERSHIP ===
divider('2', 'The bank partnership', ['Why a bank holds the money', 'Roles: the system and the machine',
                                      'Mashreq Bank Pakistan and its fit', 'The money route, collection and fees',
                                      'What Mashreq gains'])

# ---- why a bank
sl = slide('Why a Bank',
           'A company that takes money from the public and holds it is taking deposits, and only a bank may do that. '
           'So the partner bank holds and moves every rupee, and Halqa holds none.')
tb(sl, M, 1.95, CW, 0.9, [
    ('“no company shall invite, accept or renew deposits from the public: Provided that nothing in this '
     'sub-section shall apply to a banking company”', dict(size=13, font=DISPLAY, color=INK, italic=True, gap=3)),
    ('Companies Act 2017, section 84(1)', dict(size=9.5, color=GREY))])
bx = [(M, 'A company holds the money', C('F7F7F5'), GREY,
       ['Members pay into the company’s own account', 'The pooled money waits there until payout',
        'The company keeps any return on it', 'The company carries every default',
        'This is how Oraan works today']),
      (M + 6.2, 'A bank holds the money', LIME_LT, L700,
       ['Members pay from their own Mashreq accounts', 'The circle’s money sits in a Mashreq account',
        'Profit on balances goes to members as points', 'Defaults are covered by the option Mashreq chooses',
        'Halqa runs the circle and sends instructions'])]
for x, t, col, tc, pts in bx:
    rect(sl, x, 3.05, 5.9, 3.45, col)
    (icon_bank if 'bank' in t else icon_doc)(sl, x + 0.3, 3.3, 0.55, tc)
    tb(sl, x + 1.05, 3.35, 4.6, 0.45, t, size=17, color=tc, font=DISPLAY)
    tb(sl, x + 0.3, 4.05, 5.3, 2.4, pts, size=12, bullet=True, gap=6, bcolor=tc)
line(sl, M + 5.95, 4.75, M + 6.15, 4.75, L700, 1.8, arrow=True)

# ---- roles
sl = slide('Roles', 'Halqa is the system. The bank is the machine.')
colw = 5.1
lx, rx = M, W - M - colw
rect(sl, lx, 1.85, colw, 0.75, LIME, pad=(0.95, 0.1, 0.2, 0.1), anchor=MIDDLE,
     paras=[('Halqa: the system', dict(size=19, font=DISPLAY, color=INK))])
sl.shapes.add_picture(MARK, E(lx + 0.18), E(1.93), height=E(0.58))
rect(sl, rx, 1.85, colw, 0.75, LIME_LT, pad=(0.95, 0.1, 0.2, 0.1), anchor=MIDDLE,
     paras=[('Mashreq: the machine', dict(size=19, font=DISPLAY, color=L700))])
icon_bank(sl, rx + 0.2, 1.95, 0.55, L700)
tb(sl, lx + 0.05, 2.85, colw - 0.1, 3.4, [
    'Circles, rules and the order of collection', 'Credit bands and seat bands on every seat',
    'Verification and affordability models', 'The application, reminders and recovery',
    'Points, marketplace and turn market', 'Payment records for credit reporting'],
   size=13.5, bullet=True, gap=11, ind=0.3)
tb(sl, rx + 0.05, 2.85, colw - 0.1, 3.4, [
    'Accounts, account opening and customer checks', 'Direct debit mandates, collection and payouts',
    'Short term savings and profit on balances', 'Cover for defaults, in the form it chooses',
    'Credit reporting and asset financing', 'Money laundering monitoring and Shariah approval'],
   size=13.5, bullet=True, gap=11, ind=0.3)
ax1, ax2 = lx + colw + 0.2, rx - 0.2
tb(sl, ax1, 3.25, ax2 - ax1, 0.3, 'Instructions', size=10.5, color=GREY, align=CEN)
line(sl, ax1, 3.6, ax2, 3.6, L700, 1.8, arrow=True)
tb(sl, ax1, 4.3, ax2 - ax1, 0.3, 'Confirmations', size=10.5, color=GREY, align=CEN)
line(sl, ax2, 4.65, ax1, 4.65, L700, 1.8, arrow=True)
rect(sl, M, 6.05, CW, 0.62, LIME_XLT, pad=0.22, anchor=MIDDLE, paras=[
    ('Halqa never holds member money. Every instruction it sends is carried out, recorded and confirmed by '
     'Mashreq.', dict(size=12))])

# ---- mashreq
sl = slide('Mashreq Bank Pakistan',
           'A digital retail bank with an Islamic window, fully licensed since September 2025. It is gathering deposits '
           'and customers, has not yet started lending, and wins business customers through partnerships.')
events = [('19 Dec 2024', 'Restricted licence'), ('31 Jan 2025', 'Pilot operations'),
          ('29 Aug 2025', 'Islamic window approved'), ('15 Sep 2025', 'Digital retail bank licence'),
          ('16 Sep 2025', 'Commercial launch'), ('Nov 2025', 'Mashreq NEO launched')]
ty = 2.55
line(sl, M, ty, W - M, ty, L700, 1.5)
seg = CW / len(events)
for i, (d, t) in enumerate(events):
    x = M + seg * i + 0.05
    rect(sl, x, ty - 0.11, 0.22, 0.22, LIME, line=L700)
    tb(sl, x, ty - 0.55, seg - 0.1, 0.3, d, size=11, color=L700, bold=True)
    tb(sl, x, ty + 0.22, seg - 0.15, 0.6, t, size=11, spacing=1.05)
figs = [('Rs 8.5 billion', 'customer deposits, over'), ('350,000', 'customers, more than'),
        ('None', 'advances on the books'), ('9,000+', 'accounts for non resident Pakistanis, over Rs 265 million')]
fw = CW / 4
for i, (n_, t_) in enumerate(figs):
    x = M + i * fw
    rect(sl, x, 3.6, fw - 0.25, 1.35, LIME_XLT, pad=0.18, paras=[
        (n_, dict(size=24, font=DISPLAY, color=L700, gap=2)), (t_, dict(size=11, spacing=1.05))])
tb(sl, M, 5.2, CW, 1.3, [
    'Figures at 30 June 2026. The half year loss was Rs 4.7 billion; accumulated losses were Rs 16.2 billion.',
    'Mashreq NEO aims to serve 10 million Pakistanis in five years: salaried professionals, freelancers, women '
    'entrepreneurs and overseas Pakistanis. Chief executive: Muhammad Hamayun Sajjad. Chairman: Fernando Morillo.',
    'Mashreq Bank Pakistan is a member of TASDEEQ, the credit bureau.'], size=11.5, bullet=True, gap=5,
   spacing=1.08)
note(sl, 'Mashreq Bank Pakistan Limited, Directors’ Review, half year ended 30 June 2026; Mashreq press releases, '
         'January, February and 24 November 2025; Business Recorder, 16 September 2025; tasdeeq.com/members.')

# ---- fit
sl = slide('Fit with Mashreq', 'Each of Mashreq’s published aims has a direct answer in what committees bring.')
pairs = [('Serve 10 million Pakistanis in five years', 'Accounts opened in groups, a whole circle at a time'),
         ('Salaried people, freelancers, women and overseas Pakistanis', 'The same people already save through '
                                                                         'committees'),
         ('Islamic first digital banking', 'No interest, one flat fee, and a takaful option'),
         ('Accounts for non resident Pakistanis from the UAE', 'Circles joining families in the UAE and Pakistan'),
         ('No lending book yet', 'Payment records that show who repays, starting with asset finance'),
         ('Growth through partnerships', 'A partnership that brings retail customers the same way')]
lw_, gap_ = 5.2, 1.6
tb(sl, M, 1.85, lw_, 0.3, 'Mashreq’s aim or position', size=11, color=GREY, bold=True)
tb(sl, M + lw_ + gap_, 1.85, lw_, 0.3, 'What committees bring', size=11, color=GREY, bold=True)
for i, (a_, b_) in enumerate(pairs):
    y = 2.25 + i * 0.72
    rect(sl, M, y, lw_, 0.58, LIME_XLT, pad=0.16, anchor=MIDDLE, paras=[(a_, dict(size=12))])
    rect(sl, M + lw_ + gap_, y, CW - lw_ - gap_, 0.58, LIME_LT, pad=0.16, anchor=MIDDLE,
         paras=[(b_, dict(size=12, color=INK))])
    line(sl, M + lw_ + 0.05, y + 0.29, M + lw_ + gap_ - 0.05, y + 0.29, L700, 1.6, arrow=True)
note(sl, 'Mashreq, press release, 24 November 2025; Directors’ Review, half year ended 30 June 2026: business '
         'customers acquired “primarily through strategic partnerships”.')

# ---- money route
sl = slide('Money Route', 'One month of a monthly circle. Every rupee stays inside Mashreq.')
flow(sl, 1.95, 2.45, [
    ('Salary', 'Payday. Salary arrives in the member’s Mashreq account, or money is moved in from any bank.'),
    ('Savings', 'To the 8th. The instalment waits in a short term savings product for about seven days and earns '
                'profit.'),
    ('Debit', 'The 8th. Mashreq debits the instalment under the member’s mandate: contribution, cover and fee.'),
    ('Pot', 'The 8th. Contributions gather in the circle’s account at Mashreq.'),
    ('Payout', 'Mashreq pays the pot into the collecting member’s account, less any arrears owed.'),
    ('Reward', 'End of circle. Profit on the circle’s balances is credited as points, one point for each rupee.')],
    size=12, title_size=15)
rect(sl, M, 4.75, CW, 1.15, LIME_LT, pad=0.25, anchor=MIDDLE, paras=[
    ('**Halqa’s part.** Halqa works out who pays what and when, sends every instruction to Mashreq, keeps the '
     'circle’s ledger and checks it against Mashreq’s statement each morning. It never holds member money.',
     dict(size=12.5, spacing=1.1))])

# ---- collection
sl = slide('Collection',
           'A direct debit mandate at Mashreq comes first. The three payment routes already designed stay behind it.')
ladder = [('Mashreq direct debit mandate', 'The member authorises Mashreq once, inside the Halqa application. Mashreq '
           'debits the member’s account on the due date.', 'New. Default for every member'),
          ('One approval payment', 'Card, wallet or Raast Request to Pay through Safepay, approved in one step.',
           'Route 1'),
          ('Saved card auto debit', 'A charge on a card saved with Safepay. Halqa stores no card data.', 'Route 2'),
          ('Manual payment', 'Raast or bank transfer into the member’s Mashreq account from any bank or wallet.',
           'Route 3. Always available')]
lw_ = 7.75
for i, (t, b, tag) in enumerate(ladder):
    y = 1.95 + i * 1.1
    rect(sl, M + i * 0.25, y, lw_ - i * 0.25, 0.98, LIME_LT if i == 0 else LIME_XLT,
         pad=(0.85, 0.1, 2.0, 0.08), paras=[(t, dict(size=15, font=DISPLAY, color=L700, gap=3)),
                                             (b, dict(size=11.5, spacing=1.08))])
    rect(sl, M + i * 0.25, y, 0.62, 0.98, LIME, pad=0, anchor=MIDDLE,
         paras=[(str(i + 1), dict(size=20, font=DISPLAY, color=INK, align=CEN))])
    tb(sl, M + lw_ - 1.95, y + 0.12, 1.8, 0.5, tag, size=9.5, color=L700, bold=True, caps=True, track=0.5,
       align=R, spacing=1.05)
rect(sl, M + lw_ + 0.35, 1.95, CW - lw_ - 0.35, 4.3, LIME_XLT, pad=0.26, paras=[
    ('Rules', dict(size=18, font=DISPLAY, color=L700, gap=8)),
    ('Attempts on the due date, then once a day for up to five days; collection stops at the first success.',
     dict(size=12.5, bullet=True, gap=8, spacing=1.08)),
    ('No debit exceeds one instalment.', dict(size=12.5, bullet=True, gap=8)),
    ('Cancelling a mandate returns the member to manual payment. The commitment to the circle stays.',
     dict(size=12.5, bullet=True, gap=8, spacing=1.08)),
    ('Every route settles into the circle’s account at Mashreq.', dict(size=12.5, bullet=True, spacing=1.08))])
note(sl, 'Safepay holds a payment service provider licence from the State Bank of Pakistan (April 2025). No agreement '
         'is signed with Safepay or with Mashreq.')

# ---- fees
sl = slide('Fees and Income',
           'Mashreq collects one flat fee from members and shares its income from committee accounts with Halqa. '
           'The more Halqa earns from that share, the lower the member fee can go.')
bw_, by = 2.0, 2.05
xs = [M, (W - bw_) / 2, W - M - bw_]
for x, t in zip(xs, ('Member', 'Mashreq', 'Halqa')):
    rect(sl, x, by, bw_, 0.9, LIME if t != 'Mashreq' else LIME_LT, pad=0.1, anchor=MIDDLE,
         paras=[(t, dict(size=18, font=DISPLAY, color=INK if t != 'Mashreq' else L700, align=CEN))])
for a_, b_, t in [(xs[0] + bw_, xs[1], 'Instalment: contribution, cover, fee'),
                  (xs[1] + bw_, xs[2], 'Agreed share of income, each month')]:
    tb(sl, a_ + 0.1, by + 0.08, b_ - a_ - 0.2, 0.3, t, size=10.5, color=GREY, align=CEN)
    line(sl, a_ + 0.12, by + 0.45, b_ - 0.12, by + 0.45, L700, 1.8, arrow=True)
tb(sl, M, 3.35, 6.55, 3.0, [
    'One flat fee per instalment, the same for every seat, never priced by the payout month',
    'Mashreq collects the fee with each instalment as its own income',
    'Halqa receives an agreed share of Mashreq’s income from committee accounts, paid monthly',
    'As Halqa’s share rises, the member fee falls. The cost of cover appears as its own line'],
   size=13, bullet=True, gap=11, spacing=1.08, ind=0.3)
rx_ = M + 6.95
rect(sl, rx_, 3.35, W - M - rx_, 1.35, LIME_LT, pad=0.22, paras=[
    ('What a share is worth', dict(size=15, font=DISPLAY, color=L700, gap=4)),
    ('Each 10 per cent share of the bank’s net margin is about Rs 8 a member a month, on the assumptions on page 17: '
     'Rs 10 million a year at 100,000 members.', dict(size=11.5, spacing=1.08))])
rect(sl, rx_, 4.9, W - M - rx_, 1.35, LIME_XLT, pad=0.22, paras=[
    ('Against Oraan', dict(size=15, font=DISPLAY, color=L700, gap=4)),
    ('Oraan charges the first seat of a ten month committee up to 21 per cent of the pot, about 54 per cent a year on '
     'Halqa’s calculation. Halqa never prices by month.', dict(size=11.5, spacing=1.08))])

# ---- gains
sl = slide('What Mashreq Gains',
           'On stated assumptions, 100,000 members would add about Rs 2.5 billion of deposits and 100,000 accounts, '
           'against the Rs 8.5 billion Mashreq held at 30 June 2026.')
tb(sl, M, 1.95, 6.0, 0.35, 'Deposits, Rs billion', size=15, color=L700, font=DISPLAY)
hbars(sl, M, 2.5, 6.0, [('Mashreq, 30 June 2026', 8.5, 'over 8.5', L700), ('10,000 members', 0.25, '0.25', LIME),
                        ('100,000 members', 2.5, '2.5', LIME), ('1,000,000 members', 25, '25', LIME)],
      bar_h=0.42, gap=0.7, label_w=1.95, val_w=0.95, maxv=25, size=11.5)
cols = [('10,000', 'Rs 80m', 'Rs 250m', 'Rs 10m'), ('100,000', 'Rs 800m', 'Rs 2.5bn', 'Rs 100m'),
        ('1,000,000', 'Rs 8bn', 'Rs 25bn', 'Rs 1bn')]
x0, cw_ = 7.0, (W - M - 7.0 - 0.3) / 3
labels = ['members', 'committee money a month', 'average balances', 'bank’s net margin a year']
for i, c in enumerate(cols):
    x = x0 + i * (cw_ + 0.15)
    rect(sl, x, 1.95, cw_, 3.25, LIME_XLT if i != 1 else LIME_LT, pad=0.16, paras=[
        (c[0], dict(size=20, font=DISPLAY, color=L700)), (labels[0], dict(size=10, color=GREY, gap=8)),
        (c[1], dict(size=15, bold=True)), (labels[1], dict(size=10, color=GREY, gap=8)),
        (c[2], dict(size=15, bold=True)), (labels[2], dict(size=10, color=GREY, gap=8)),
        (c[3], dict(size=15, bold=True)), (labels[3], dict(size=10, color=GREY))])
tb(sl, x0, 5.4, W - M - x0, 1.1, [
    'Also: women savers, customers arriving in groups of 6 to 20, an Islamic product, a route to overseas '
    'Pakistanis, and fee income from cover and financing.'], size=11.5, spacing=1.1)
note(sl, 'Assumptions, to be replaced by Mashreq’s own: average instalment Rs 8,000 a month; average balance Rs 25,000 '
         'a member once salaries move to Mashreq; net margin to the bank of 4 per cent a year after the profit credited '
         'to members.')

# ============================================================ 3 PRODUCT ===
divider('3', 'The product', ['From signup to a completed circle', 'Checks, bands and cover',
                             'Late payment, default and exit', 'Credit reporting, points and the turn market',
                             'Products in stages, and Hyper'])

# ---- signup to completion
sl = slide('From Signup to Completion', 'Each step carries the control that applies to it.')
steps = [('Sign up', ['Phone number and one time passcode', 'App PIN on every open', 'Name, address, city and occupation']),
         ('Verify', ['Mashreq account opening with its customer checks', 'CNIC and face match',
                     'Halqa’s income account check']),
         ('Join', ['Every amount shown before signing', 'Seat chosen within the band', 'Affordability check',
                   'Undertaking signed']),
         ('Pay', ['Mashreq direct debit on the 8th', 'Safepay, Raast or transfer as back up',
                  'Receipt showing each part']),
         ('Collect', ['Pot paid into the member’s Mashreq account', 'Paid once and only once',
                      'Remaining instalments restated']),
         ('Complete', ['Every payment reported to the bureau', 'Band rises on a clean circle',
                       'Profit credited as points'])]
bw = (CW - 5 * 0.25) / 6
for i, (t, pts) in enumerate(steps):
    x = M + i * (bw + 0.25)
    rect(sl, x, 1.9, bw, 0.6, LIME if i in (0, 5) else LIME_LT, pad=0.08, anchor=MIDDLE,
         paras=[(t, dict(size=14, bold=True, color=INK, align=CEN))])
    if i < 5:
        line(sl, x + bw + 0.02, 2.2, x + bw + 0.23, 2.2, L700, 1.5, arrow=True)
    tb(sl, x + 0.02, 2.75, bw - 0.04, 3.0, pts, size=11, bullet=True, gap=6, spacing=1.05)
rect(sl, M, 5.55, CW, 0.95, LIME_XLT, pad=0.22, anchor=MIDDLE, paras=[
    ('**Messages and records.** Notices on the WhatsApp Business Platform, one daily notice on Hyper, and statements, '
     'receipts, the schedule and a full data export for every member.', dict(size=12, spacing=1.08))])

# ---- checks and bands
sl = slide('Checks and Bands',
           'Mashreq checks who the customer is. Halqa checks whether the member can carry the circle, and the bands '
           'decide which seats are open. Both apply to every seat, with no exception.')
checks = [('Identity', 'Mashreq’s account opening, CNIC and face match'),
          ('Income account', 'In the member’s name, with a regular salary pattern'),
          ('Affordability', 'All instalments within a third of verified income; 40 per cent with other loans'),
          ('Credit report', 'TASDEEQ report on the member’s authority'),
          ('Seat', 'Offered only within the member’s band')]
for i, (t, d) in enumerate(checks):
    y = 2.05 + i * 0.86
    wd = 6.0 - i * 0.35
    rect(sl, M, y, wd, 0.72, LIME_LT if i % 2 == 0 else LIME_XLT, pad=0.16, anchor=MIDDLE, paras=[
        ([('%s   ' % t, dict(size=13, bold=True, color=L700)), (d, dict(size=11.5))], {})])
x0 = 7.3
tb(sl, x0, 2.0, 5.4, 0.3, 'Bands on the Halqa score, 300 to 850', size=13, color=L700, font=DISPLAY)
bands = [('Excellent', '750 and above', 'Any seat', L700, WHITE), ('Good', '650 to 749', 'Any seat', L600, WHITE),
         ('Fair', '550 to 649', 'Second half', LIME, INK), ('Rebuilding', 'below 550', 'Last three', LIME_LT, INK),
         ('New member', 'any score', 'Last three, until two clean circles and a call', LIME_XLT, INK)]
for i, (n_, sc, st_, col, tc) in enumerate(bands):
    y = 2.45 + i * 0.75
    rect(sl, x0, y, 1.6, 0.62, col, pad=0.1, anchor=MIDDLE, paras=[(n_, dict(size=12, bold=True, color=tc,
                                                                                align=CEN))])
    tb(sl, x0 + 1.75, y + 0.04, 1.3, 0.55, sc, size=11, color=GREY, anchor=MIDDLE)
    tb(sl, x0 + 3.05, y + 0.04, W - M - x0 - 3.05, 0.6, st_, size=11.5, anchor=MIDDLE, spacing=1.05)
note(sl, 'Rebuilding members may not buy turns. A TASDEEQ score of 200 to 600 maps onto the Halqa scale in a straight '
         'line. Source: score-bands.ts; the maths documents HQ-MF-01 to HQ-MF-06.')

# ---- cover
sl = slide('Cover',
           'Three ways to protect the other members when someone collects early and stops paying. Mashreq chooses one.')
opts = [('Bank guarantee', 'Mashreq guarantees each pot from its balance sheet and prices the risk into the fee. '
                           'Simplest for members; uses the bank’s capital.', 'Mashreq', 'The member fee',
         'Subject to the Shariah board'),
        ('Insurance', 'A policy on each circle through the bank’s own insurance arrangements. A premium is part of '
                      'each instalment.', 'The insurer', 'A premium in each instalment', 'Conventional only'),
        ('Takaful', 'Cover from a general takaful operator for the Islamic window. A contribution is part of each '
                    'instalment.', 'The takaful fund', 'A contribution in each instalment', 'Yes')]
cw_ = (CW - 2 * 0.3) / 3
for i, (t, d, who, paid, isl) in enumerate(opts):
    x = M + i * (cw_ + 0.3)
    (icon_bank, icon_doc, icon_group)[i](sl, x + 0.25, 2.1, 0.5, L700)
    tb(sl, x + 0.95, 2.12, cw_ - 1.0, 0.5, t, size=19, color=L700, font=DISPLAY)
    tb(sl, x + 0.25, 2.8, cw_ - 0.4, 1.3, d, size=11.5, spacing=1.1)
    for k, (lab, val) in enumerate((('Carries the loss', who), ('Paid through', paid), ('Islamic window', isl))):
        y = 4.2 + k * 0.62
        rect(sl, x, y, cw_, 0.54, LIME_XLT if k % 2 == 0 else LIME_LT, pad=0.14, anchor=MIDDLE, paras=[
            ([(lab + '   ', dict(size=10, color=GREY)), (val, dict(size=11.5, bold=True))], {})])
rect(sl, M, 6.2, CW, 0.5, LIME, pad=0.2, anchor=MIDDLE, paras=[
    ('In every option: mandatory bands, arrears taken from the late member’s own pot, the late ladder, and a hard stop '
     'on any circle whose unrecovered exposure exceeds its cover.', dict(size=11.5, color=INK))])

# ---- default
sl = slide('Late Payment, Default and Exit',
           'Most members never reach recovery. When they do, each step is stated in advance, and nothing is done by '
           'harassment.')
tl = [('Due date', 'The 8th. Debit attempted'), ('Grace', 'Reminders, retries for up to five days'),
      ('Late ladder', 'Three steps, each with a stated penalty'), ('Recovery case', 'Hardship statement or a revised '
                                                                                 'date'),
      ('Pot set off', 'Arrears taken from the member’s own pot'), ('Cover claim', 'Guarantee, insurer or takaful fund'),
      ('Civil suit', 'Last step, on the signed undertaking')]
ty = 2.35
line(sl, M, ty, W - M, ty, L700, 2.0, arrow=True)
seg = (CW - 0.2) / len(tl)
for i, (t, d) in enumerate(tl):
    x = M + i * seg
    rect(sl, x, ty - 0.1, 0.2, 0.2, LIME if i < 4 else L700)
    tb(sl, x, ty - 0.55, seg - 0.05, 0.35, t, size=12, bold=True, color=L700)
    tb(sl, x, ty + 0.2, seg - 0.12, 0.8, d, size=10.5, spacing=1.05)
tb(sl, M, 3.55, 5.9, 0.35, 'Exit ladder', size=15, color=L700, font=DISPLAY)
ex = ['Withdraw inside the 24 hour window: free and unrecorded', 'Substitution: a replacement takes the seat',
      'Group approved exit: a 72 hour vote', 'Hardship exit: recorded as hardship, not default',
      'Abandonment: not an exit; recovery applies']
for i, t in enumerate(ex):
    rect(sl, M + i * 0.35, 4.0 + i * 0.5, 5.9 - i * 0.35, 0.42, LIME_LT if i < 4 else C('F7F7F5'), pad=0.12,
         anchor=MIDDLE, paras=[('%d   %s' % (i + 1, t), dict(size=11))])
x0 = 7.0
rect(sl, x0, 3.55, W - M - x0, 2.95, LIME_XLT, pad=0.24, paras=[
    ('What is never done', dict(size=15, font=DISPLAY, color=L700, gap=6)),
    ('No collection calls and no use of contact lists, at any stage', dict(size=12, bullet=True, gap=5)),
    ('No public list of defaulters; the flag stays internal', dict(size=12, bullet=True, gap=5)),
    ('On Shariah labelled circles, penalties go to charity at the end of the circle', dict(size=12, bullet=True,
                                                                                            spacing=1.05))])
note(sl, 'A summary suit lies only on a cheque or other negotiable instrument (Code of Civil Procedure, Order 37). '
         'Source: exit-ladder.ts; Default Prevention (HQ-CP-05).')

# ---- credit
sl = slide('Credit Reporting and Financing',
           'Every payment is reported from the first day, good and bad. Mashreq is already a member of TASDEEQ, which '
           'gives a direct route for committee records, subject to TASDEEQ accepting the data.')
nodes = [('Member pays', 'on time or late'), ('Halqa records', 'each payment, each circle'),
         ('Reported', 'by Mashreq through its TASDEEQ membership, or to TASDEEQ directly'),
         ('Credit record', 'readable by any lender')]
bw = (CW - 3 * 0.35) / 4
for i, (t, d) in enumerate(nodes):
    x = M + i * (bw + 0.35)
    rect(sl, x, 1.95, bw, 1.15, LIME_LT if i == 2 else LIME_XLT, pad=0.16, paras=[
        (t, dict(size=15, font=DISPLAY, color=L700, gap=3)), (d, dict(size=11, spacing=1.05))])
    if i < 3:
        line(sl, x + bw + 0.04, 2.5, x + bw + 0.31, 2.5, L700, 1.6, arrow=True)
tb(sl, M, 3.45, 6.0, 0.35, 'Financing that follows', size=15, color=L700, font=DISPLAY)
fin = ['Circle completed on time', 'Record reported', 'Mashreq offers finance for a motorcycle or an appliance',
       'Asset circles financed by Mashreq']
for i, t in enumerate(fin):
    y = 3.95 + i * 0.62
    rect(sl, M, y, 6.0, 0.5, LIME if i == 3 else LIME_XLT, pad=0.16, anchor=MIDDLE, paras=[(t, dict(size=12))])
x0 = 7.0
rect(sl, x0, 3.45, W - M - x0, 3.0, LIME_XLT, pad=0.24, paras=[
    ('Evidence from elsewhere', dict(size=15, font=DISPLAY, color=L700, gap=6)),
    ('Mission Asset Fund, San Francisco: lending circles reported to the three US bureaus; average score increase of '
     '168 points.', dict(size=11.5, bullet=True, gap=6, spacing=1.08)),
    ('Esusu, United States: rent reported to the bureaus; valued at US$1 billion in January 2022.',
     dict(size=11.5, bullet=True, gap=6, spacing=1.08)),
    ('In Pakistan the payment households never miss is the committee instalment, ranked above rent.',
     dict(size=11.5, bullet=True, spacing=1.08))])
note(sl, 'TASDEEQ members list (tasdeeq.com/members, read 28 September 2026); Credit Bureaus Act 2015. Mashreq '
         'reported no advances at 30 June 2026.')

# ---- points
sl = slide('Points and Marketplace',
           'Profit earned on committee balances comes back to members at the end of each circle as points, one point '
           'for each rupee.')
cx, cy = W / 2, 3.95
circ = [('Earn', ['Profit on the circle’s balances', 'On time payments and completed circles', 'Referrals',
                  'Points bought with money'], M + 0.2),
        ('Hold', ['A points balance issued under Mashreq’s licence', 'One point equals one rupee',
                  'No cash withdrawal'], M + 4.35),
        ('Spend', ['Goods in the Halqa marketplace', 'Halqa fees', 'Turns in the turn market',
                   'Exchange with the bank’s own rewards'], M + 8.5)]
for i, (t, pts, x) in enumerate(circ):
    rect(sl, x, 1.95, 3.5, 2.85, LIME_XLT if i != 1 else LIME_LT, pad=0.22, paras=[
        (t, dict(size=19, font=DISPLAY, color=L700, gap=6))] +
         [(p_, dict(size=12, bullet=True, gap=5, spacing=1.05)) for p_ in pts])
    if i < 2:
        line(sl, x + 3.55, 3.35, x + 4.1, 3.35, L700, 1.8, arrow=True)
line(sl, M + 10.2, 4.9, M + 10.2, 5.2, L700, 1.2)
line(sl, M + 10.2, 5.2, M + 1.9, 5.2, L700, 1.2)
line(sl, M + 1.9, 5.2, M + 1.9, 4.85, L700, 1.2, arrow=True)
tb(sl, M + 3.5, 5.24, 5.0, 0.3, 'spent points bring members back to the next circle', size=10, color=GREY,
   align=CEN)
rect(sl, M, 5.65, 5.9, 0.85, LIME_LT, pad=0.18, paras=[
    ('**Online marketplaces.** Points redeemed at checkout or for vouchers with marketplaces in Pakistan, such as '
     'Daraz. None contracted.', dict(size=11, spacing=1.05))])
rect(sl, M + 6.19, 5.65, CW - 6.19, 0.85, LIME_LT, pad=0.18, paras=[
    ('**Bank rewards.** Points exchanged with Mashreq’s own rewards programme, where one runs, at an agreed rate.',
     dict(size=11, spacing=1.05))])
note(sl, 'Points that can be bought, transferred and spent outside Halqa are stored value, so they run as Mashreq’s '
         'product under its licence and its approval.')

# ---- turn market
sl = slide('Turn Market', 'Members can buy and sell turns with each other, inside the bands.')
seats = 12
sw_ = (CW - 11 * 0.08) / seats
for k in range(seats):
    x = M + k * (sw_ + 0.08)
    col = LIME if k in (2, 10) else LIME_XLT
    rect(sl, x, 2.05, sw_, 0.62, col, line=L700 if k in (2, 10) else RULE, pad=0, anchor=MIDDLE,
         paras=[('%d' % (k + 1), dict(size=13, bold=k in (2, 10), align=CEN))])
xa = M + 2 * (sw_ + 0.08) + sw_ / 2
xb = M + 10 * (sw_ + 0.08) + sw_ / 2
line(sl, xa, 2.72, xa, 3.15, L700, 1.4)
line(sl, xa, 3.15, xb, 3.15, L700, 1.4)
line(sl, xb, 3.15, xb, 2.75, L700, 1.4, arrow=True)
line(sl, xb, 1.98, xb, 1.72, L700, 1.4)
line(sl, xb, 1.72, xa, 1.72, L700, 1.4)
line(sl, xa, 1.72, xa, 1.99, L700, 1.4, arrow=True)
tb(sl, (xa + xb) / 2 - 3.0, 3.2, 6.0, 0.3, 'The buyer moves from seat 11 to seat 3 and pays 40,000 points',
   size=11, color=L700, align=CEN, bold=True)
flow(sl, 3.85, 1.55, [('List', 'Asking price in points or rupees'), ('Offer', 'The buyer’s band must permit the seat'),
                      ('Accept', 'The seller accepts with a PIN'), ('Approve', 'The host approves'),
                      ('Settle', 'Seats swap; the price moves between accounts at Mashreq'),
                      ('Record', 'The price goes on the ledger')], size=10.5, title_size=14, gap=0.3)
rect(sl, M, 5.65, CW, 0.9, C('F7F7F5'), pad=0.2, anchor=MIDDLE, paras=[
    ('Not on Hyper or asset circles. A price ceiling for each seat and an exchange fee paid by the buyer. The product '
     'needs Mashreq’s approval and a Shariah board ruling, and counsel must confirm whether a turn sold for money is '
     'lending between members. Circles that allow turn trades are not labelled Shariah compliant.',
     dict(size=11, spacing=1.08))])

# ---- products and stages
sl = slide('Products and Stages',
           'Monthly circles come first. Each later product opens after the pilot, with Mashreq’s approval.')
stages = [('Pilot', LIME, [('Monthly circles', 'Members who know each other, 6 to 20 in a circle'),
                            ('Short term savings', 'The instalment earns profit from payday to the 8th')]),
          ('After the pilot', LIME_LT, [('Open circles', 'Members matched by band, with cover'),
                                        ('Asset circles', 'A motorcycle or appliance, financed by Mashreq'),
                                        ('Across borders', 'Members in the UAE pay from accounts for non resident '
                                                           'Pakistanis')]),
          ('Experimental', LIME_XLT, [('Hyper', 'Daily circles in two options, labelled experimental')])]
x = M
unit = (CW - 2 * 0.2) / 6
for t, col, prods in stages:
    wd = unit * len(prods)
    rect(sl, x, 1.95, wd, 0.55, col, pad=0.12, anchor=MIDDLE, paras=[(t, dict(size=14, bold=True, color=INK))])
    for k, (pn, pd) in enumerate(prods):
        px = x + k * (wd / len(prods))
        rect(sl, px, 2.62, wd / len(prods) - 0.12, 2.1, LIME_XLT, line=RULE, pad=0.16, paras=[
            (pn, dict(size=15, font=DISPLAY, color=L700, gap=4)), (pd, dict(size=11.5, spacing=1.08))])
    x += wd + 0.2
tb(sl, M, 5.0, CW, 1.4, [
    'Each product keeps the same rules: one flat fee, bands on every seat, cover chosen by Mashreq, and every payment '
    'reported.',
    'Asset circles replace the modaraba design of August: Mashreq finances the asset, and the circle runs as usual.'],
   size=12, bullet=True, gap=6, spacing=1.08)

# ---- hyper
sl = slide('Hyper', 'Daily circles for members with daily income, kept exactly as designed and labelled experimental.')
rect(sl, M + 1.45, 0.66, 1.35, 0.36, AMBER, pad=0.05, anchor=MIDDLE,
     paras=[('Experimental', dict(size=10.5, bold=True, color=INK, align=CEN))])
optsh = [('Option 1', [('400', 'members'), ('50', 'days'), ('Rs 450', 'paid each day'), ('Rs 15,000', 'pot, '
                                                                                                    'collected once'),
                       ('8', 'collect each day')]),
         ('Option 2', [('390', 'members'), ('26', 'days, Sundays off'), ('Rs 500', 'paid each day'),
                       ('Rs 8,666.67', 'pot, collected once'), ('15', 'collect each day')])]
for i, (t, vals) in enumerate(optsh):
    y = 1.95 + i * 1.9
    tb(sl, M, y, 1.5, 0.4, t, size=17, color=L700, font=DISPLAY)
    vw = (CW - 1.6) / 5
    for k, (n_, l_) in enumerate(vals):
        x = M + 1.6 + k * vw
        rect(sl, x, y - 0.05, vw - 0.12, 1.55, LIME_XLT if i == 0 else LIME_LT, pad=0.14, paras=[
            (n_, dict(size=21, font=DISPLAY, color=L700, gap=2)), (l_, dict(size=10.5, spacing=1.05))])
tb(sl, M, 5.85, CW, 0.8, [
    ('Each daily payment is three parts: the contribution to the pot (Rs 300 or Rs 333.33), cover (Rs 75 or Rs 83.33) '
     'and the fee (Rs 75 or Rs 83.33). pot = contribution × days, and roster = collectors each day × days.',
     dict(size=11.5, spacing=1.08))])
note(sl, 'Source: Hyper Committee (HQ-CP-08). Entry needs a linked income account, two clean circles, the highest '
         'identity level and daily income of Rs 1,000 or more on five days of every week for eight weeks.')

# ---- hyper day and stress
sl = slide('Hyper: One Day, and the Stress Bands',
           'On Option 1, 392 members pay each day and 8 collect. Each collector is paid by 49 named members, so the '
           'daily sums divide exactly.')
tb(sl, M, 1.95, 6.0, 0.3, 'One day on Option 1', size=14, color=L700, font=DISPLAY)
for k in range(8):
    x = M + 0.1 + k * 0.72
    rect(sl, x, 2.45, 0.6, 0.6, LIME, pad=0, anchor=MIDDLE, paras=[('C%d' % (k + 1), dict(size=10, bold=True,
                                                                                          align=CEN))])
    for j in range(7):
        rect(sl, x + (j % 4) * 0.15, 3.25 + (j // 4) * 0.17, 0.11, 0.11, L700)
    line(sl, x + 0.3, 3.2, x + 0.3, 3.08, L700, 1.0, arrow=True)
tb(sl, M, 3.75, 6.0, 1.4, [
    '49 payers × Rs 300 = Rs 14,700, plus the collector’s own Rs 300 set off = Rs 15,000',
    'Option 2: 375 payers, 15 collectors, 25 each: 25 × Rs 333.33 + Rs 333.33 = Rs 8,666.67'],
   size=11.5, bullet=True, gap=6, spacing=1.08)
x0 = 7.1
tb(sl, x0, 1.95, W - M - x0, 0.3, 'Stress index, 0 to 100, recomputed daily', size=14, color=L700, font=DISPLAY)
gw_ = W - M - x0
for col, lab, a_, b_ in ((LIME, 'Green 0 to 33', 0, 0.33), (AMBER, 'Amber 34 to 66', 0.33, 0.66),
                         (RED, 'Red 67 to 100', 0.66, 1.0)):
    rect(sl, x0 + gw_ * a_, 2.45, gw_ * (b_ - a_), 0.5, col, pad=0.05, anchor=MIDDLE,
         paras=[(lab, dict(size=10.5, bold=True, color=INK if col != RED else WHITE, align=CEN))])
acts = [('Day 5', '1.2', 'Reminders only'), ('Day 20', '28.6', 'Host told; cover provider on notice'),
        ('Day 30', '49.5', 'New joins blocked'), ('Day 45', '72.1', 'Next day not opened')]
for i, (d, v, a_) in enumerate(acts):
    y = 3.2 + i * 0.62
    tb(sl, x0, y, 0.9, 0.5, d, size=11.5, bold=True, color=L700, anchor=MIDDLE)
    tb(sl, x0 + 0.95, y, 0.8, 0.5, v, size=11.5, anchor=MIDDLE)
    tb(sl, x0 + 1.8, y, gw_ - 1.8, 0.5, a_, size=11.5, anchor=MIDDLE)
    line(sl, x0, y + 0.56, W - M, y + 0.56, RULE, 0.75)
rect(sl, M, 5.7, CW, 0.8, LIME_XLT, pad=0.2, anchor=MIDDLE, paras=[
    ('**Hard stop.** Once unrecovered exposure exceeds the cover limit, no further day opens until it is back below, '
     'whatever the index says. The index weighs severity 45, breadth 20, behaviour 15, persistence 10, timing 10.',
     dict(size=11.5, spacing=1.08))])
note(sl, 'Modelled. Example: an Option 1 circle with cover of eight pots. Source: Hyper Committee (HQ-CP-08), sections '
         '8 and 9.')

# ================================================= 4 EVIDENCE AND COMPETITION ===
divider('4', 'Evidence and competition', ['What twenty five attempts show', 'The models that work',
                                          'The competition, and Oraan in detail', 'The loan app precedent'])

# ---- 25 attempts
sl = slide('What 25 Attempts Show',
           'Every failure held member money without a licence, guaranteed it, or removed the rotation. Every lasting '
           'success kept the money outside itself or paid the full regulatory cost of holding it.')
causes = [('Holding the pool', 'Punjab cooperatives, TAG, Sidra Humaid, Saradha'),
          ('Guaranteeing payouts', 'eMoneyPool'), ('Pooling strangers', 'Yahoo Tanda, Puddle'),
          ('No early pot', 'UBL Kommittee'), ('Adding interest', 'UBL Kommittee, auction models'),
          ('A ledger without payments', 'Udhaar Book, DigiKhata'), ('Charging for access', 'Subscription '
                                                                                           'committee apps'),
          ('Rules written after a scandal', 'India after Saradha; Pakistan after the loan apps')]
tw = (CW - 3 * 0.25) / 4
for i, (t, cases) in enumerate(causes):
    rr, cc = divmod(i, 4)
    x, y = M + cc * (tw + 0.25), 2.05 + rr * 2.0
    rect(sl, x, y, tw, 1.8, LIME_XLT if (rr + cc) % 2 == 0 else LIME_LT, pad=0.18, paras=[
        ('%d' % (i + 1), dict(size=22, font=DISPLAY, color=L700, gap=2)),
        (t, dict(size=14, bold=True, gap=4)), (cases, dict(size=11, color=GREY, spacing=1.05))])
rect(sl, M, 6.1, CW, 0.52, LIME, pad=0.2, anchor=MIDDLE, paras=[
    ('With a partner bank holding the money under its licence, Halqa avoids the first cause and keeps the rotation, '
     'the peers and the flat fee.', dict(size=11.5))])

# ---- models that work
sl = slide('The Models That Work', 'Four markets where digital committees reached scale, and what each one shows.')
cases = [('Egypt', 'Money Fellows', '8 million downloads, about US$1.5 billion processed, profitable in 2025',
          'Holds the money under a licence and covers defaults from its own capital'),
         ('Saudi Arabia', 'Hakbah', '1.3 million users, central bank sandbox permit',
          'Worked with the central bank as its first partner'),
         ('Indonesia', 'Mapan', 'Bought by GO-JEK in 2017', 'Grew through village organisers paid for organising'),
         ('India', 'The Money Club', 'About 200,000 users and 17,000 clubs',
          'Members start small; a clean history opens bigger clubs')]
cw_ = (CW - 3 * 0.25) / 4
for i, (ctry, name, scale, lesson) in enumerate(cases):
    x = M + i * (cw_ + 0.25)
    rect(sl, x, 1.95, cw_, 0.55, LIME, pad=0.12, anchor=MIDDLE, paras=[(ctry, dict(size=13, bold=True))])
    rect(sl, x, 2.5, cw_, 3.6, LIME_XLT, pad=0.18, paras=[
        (name, dict(size=18, font=DISPLAY, color=L700, gap=6)), (scale, dict(size=11.5, spacing=1.08, gap=10)),
        ('What it shows', dict(size=10, color=GREY, bold=True, gap=2)), (lesson, dict(size=11.5, spacing=1.08))])
note(sl, 'Reported, from company disclosures and the press. Halqa’s route follows Hakbah’s: a regulated partner '
         'first, then scale.')

# ---- competition
sl = slide('Competition',
           'Who holds the money and how the fee is set separate every provider. Halqa with Mashreq is the only one with '
           'a bank holding the money and one flat fee for every seat.')
gx, gy, gw_, gh_ = M + 0.9, 1.95, 7.4, 4.35
rect(sl, gx, gy, gw_, gh_, LIME_XLT)
line(sl, gx, gy + gh_, gx + gw_ + 0.15, gy + gh_, INK, 1.0, arrow=True)
line(sl, gx, gy + gh_, gx, gy - 0.15, INK, 1.0, arrow=True)
for k, t in enumerate(('The organiser', 'The platform company', 'A bank')):
    tb(sl, gx + k * gw_ / 3, gy + gh_ + 0.08, gw_ / 3, 0.3, t, size=10.5, color=GREY, align=CEN)
    if k:
        line(sl, gx + k * gw_ / 3, gy, gx + k * gw_ / 3, gy + gh_, RULE, 0.75, dash=True)
tb(sl, M - 0.3, gy + 0.1, 1.1, 0.6, 'Flat or none', size=10, color=GREY, align=R)
tb(sl, M - 0.3, gy + gh_ - 0.75, 1.1, 0.6, 'Priced by payout month or interest', size=10, color=GREY, align=R)
tb(sl, gx, gy + gh_ + 0.38, gw_, 0.3, 'Who holds the money', size=10.5, color=INK, bold=True, align=CEN)
dots = [('Informal committee', 0.5, 0.18, GREY), ('JazzCash Committee', 0.5, 0.45, GREY),
        ('Oraan', 1.5, 0.85, RED), ('Money Fellows', 1.62, 0.7, GREY), ('UBL Kommittee', 2.35, 0.85, GREY),
        ('Halqa with Mashreq', 2.5, 0.12, L700)]
for n_, xf, yf, col in dots:
    px, py = gx + xf * gw_ / 3, gy + yf * gh_
    shape(sl, MSO_SHAPE.OVAL, px - 0.13, py - 0.13, 0.26, 0.26, LIME if 'Halqa' in n_ else col)
    if xf > 2.0:
        tb(sl, px - 2.48, py - 0.17, 2.3, 0.35, n_, size=11, bold='Halqa' in n_, align=R,
           color=L700 if 'Halqa' in n_ else INK)
    else:
        tb(sl, px + 0.18, py - 0.17, 2.3, 0.35, n_, size=11, bold='Halqa' in n_, color=L700 if 'Halqa' in n_ else INK)
x0 = M + 8.75
tb(sl, x0, 1.95, W - M - x0, 4.5, [
    ('JazzCash Committee', dict(size=13, bold=True, color=L700, gap=2)),
    ('Launched 13 August 2026. The pot collects in the organiser’s wallet and the organiser pays out; no exit before '
     'the end; fees not published. The biggest threat, by reach.', dict(size=10.5, spacing=1.05, gap=8)),
    ('Oraan', dict(size=13, bold=True, color=L700, gap=2)),
    ('Holds members’ money in its own accounts and prices early seats up to 21 per cent of the pot.',
     dict(size=10.5, spacing=1.05, gap=8)),
    ('UBL Kommittee', dict(size=13, bold=True, color=L700, gap=2)),
    ('A bank deposit sold under the committee name, with no rotation and a fixed return.',
     dict(size=10.5, spacing=1.05))])

# ---- oraan
sl = slide('Oraan in Detail',
           'The closest competitor. Its committee company holds the money without a financial licence, keeps the '
           'return on it and carries defaults on its own balance sheet.')
rect(sl, M + 1.3, 1.95, 3.6, 0.8, LIME_XLT, line=L700, pad=0.12, anchor=MIDDLE, paras=[
    ('ORAAN PTE. LTD., Singapore', dict(size=12, bold=True, align=CEN)),
    ('registered 20 April 2018; publishes the apps; ownership not published', dict(size=10, color=GREY, align=CEN))])
for k, (t, d, col) in enumerate((('Oraan Tech, Karachi', 'Runs the committees and holds the money. No financial '
                                                        'licence', C('F7F7F5')),
                                 ('Oraan Financial Services', 'SECP lending licence since 3 June 2024, used for '
                                                              'education loans. TASDEEQ member', LIME_LT))):
    x = M + k * 3.3
    rect(sl, x, 3.15, 3.1, 1.1, col, line=L700, pad=0.12, paras=[(t, dict(size=12, bold=True, gap=2)),
                                                                   (d, dict(size=10, spacing=1.05))])
    line(sl, M + 3.1, 2.75, x + 1.55, 3.12, L700, 1.2, arrow=True)
tb(sl, M, 4.42, 6.4, 0.3, 'Fee each month by payout month, 10 month committee', size=12, color=L700, font=DISPLAY)
fees = [21, 19, 16.5, 13.5, 10, 6, 2, 0, 0, 0]
for k, v in enumerate(fees):
    x = M + 0.1 + k * 0.62
    hgt = 1.1 * v / 21.0
    rect(sl, x, 6.3 - hgt, 0.46, max(hgt, 0.02), RED if k == 0 else GREY_LT)
    tb(sl, x - 0.08, 6.33, 0.62, 0.25, str(k + 1), size=9, color=GREY, align=CEN)
    tb(sl, x - 0.1, 6.3 - hgt - 0.26, 0.66, 0.25, ('%g%%' % v) if v else '0', size=9, color=INK, align=CEN)
x0 = 7.25
tb(sl, x0, 1.95, W - M - x0, 4.6, [
    ('Money', dict(size=13, bold=True, color=L700, gap=2)),
    ('Paid into Oraan Tech’s own bank accounts by 1LINK bill number or transfer; held as “Amanat”; any return kept '
     'by Oraan; paid out between the 11th and 18th.', dict(size=10.5, spacing=1.05, gap=7)),
    ('Default', dict(size=13, bold=True, color=L700, gap=2)),
    ('30 days unpaid after collecting: the whole balance recalled with legal costs after 30 days’ notice, a suit, a '
     'bureau report, and family payouts held back.', dict(size=10.5, spacing=1.05, gap=7)),
    ('Members', dict(size=13, bold=True, color=L700, gap=2)),
    ('The only count of paying members ever published is 10,000 in 2021. The 600,000 figure counts accounts; its own '
     'committee page says “thousands of members”.', dict(size=10.5, spacing=1.05, gap=7)),
    ('Links', dict(size=13, bold=True, color=L700, gap=2)),
    ('A Karandaaz grant funded by the Gates Foundation, 2019 and 2020; Dubai Islamic Bank for payouts in 2021; its '
     'platform licensed to an unnamed bank abroad in 2026.', dict(size=10.5, spacing=1.05))])
note(sl, 'Source: Oraan Research (HQ-RS-01), 28 September 2026, from Oraan’s terms, the SECP and Singapore registers, '
         'the TASDEEQ members list, the Karandaaz case study and the press.')

# ---- loan apps
sl = slide('The Loan App Precedent',
           'Four hundred predatory loan apps were blocked in 2023 and 2024. Every permission Halqa asks for is read '
           'against them.')
rect(sl, M, 1.95, 5.9, 2.1, C('F7F7F5'), pad=0.22, paras=[
    ('How they worked', dict(size=15, font=DISPLAY, color=GREY, gap=5)),
    ('Loans of Rs 1,000 to 25,000 over 7 to 90 days, priced for mass default and recovered by harassment: harvested '
     'contacts, calls to relatives, altered photographs. A case in Rawalpindi ended in a suicide and more than twenty '
     'arrests.', dict(size=11.5, spacing=1.1))])
rect(sl, M, 4.25, 5.9, 2.25, C('F7F7F5'), pad=0.22, paras=[
    ('What was banned, mostly even with consent', dict(size=15, font=DISPLAY, color=GREY, gap=5)),
    ('SECP circulars of 2023 and 2024: no contacts or gallery; contact only a consented guarantor; a cap of Rs 25,000 '
     'an app; a key fact statement; data kept in Pakistan. Google Play, 31 May 2023: no contacts, photos, precise '
     'location or storage.', dict(size=11.5, spacing=1.1))])
x0 = M + 6.2
tb(sl, x0, 1.95, W - M - x0, 0.4, 'What Halqa refuses', size=16, color=L700, font=DISPLAY)
refuse = ['Contact lists', 'Gallery, text messages, call log and installed apps', 'Background location',
          'Reading the screens of bank apps', 'A public defaulter list']
for i, t in enumerate(refuse):
    y = 2.5 + i * 0.62
    cross(sl, x0, y + 0.08, 0.26, RED, 2.2)
    tb(sl, x0 + 0.45, y, W - M - x0 - 0.45, 0.5, t, size=12.5, anchor=MIDDLE)
rect(sl, x0, 5.7, W - M - x0, 0.8, LIME_LT, pad=0.18, anchor=MIDDLE, paras=[
    ('Oraan’s scoring partner reads installed apps, web history and bookmarks. Halqa asks for none of these.',
     dict(size=11.5, spacing=1.05))])

# ================================================= 5 LAW, RISK AND PLAN ===
divider('5', 'Law, risk and plan', ['The legal map', 'Risks and controls', 'How the systems connect',
                                    'Where Halqa stands', 'Approvals and pilot', 'Measures and open questions'])

# ---- legal map
sl = slide('Legal Map', 'Each activity, the law that governs it, and who carries the duty under the bank route.')
rows_ = [('Holding and moving money', 'Banking Companies Ordinance; State Bank rules', 'Mashreq'),
         ('Direct debit mandates', 'Payment Systems and Electronic Fund Transfers Act 2007', 'Mashreq'),
         ('Card, wallet and Raast payments', 'PSO/PSP Rules 2014', 'Safepay'),
         ('Deposits by a company', 'Companies Act 2017, section 84', 'Avoided: Halqa holds nothing'),
         ('Halqa as a service provider', 'State Bank framework for outsourcing', 'Mashreq’s board, Halqa assessed'),
         ('Credit reporting', 'Credit Bureaus Act 2015', 'Mashreq or TASDEEQ, with consent'),
         ('Cover', 'Insurance Ordinance 2000; Takaful Rules', 'Mashreq, insurer or takaful operator'),
         ('Points bought with money', 'Electronic money rules', 'Mashreq’s product and licence'),
         ('Turn market', 'Lending and Shariah questions', 'Counsel’s opinion before launch'),
         ('Hyper fee', 'Flat service fee, never graded by day', 'Halqa, as company revenue')]
for i, (a_, b_, c_) in enumerate(rows_):
    y = 1.85 + i * 0.47
    rect(sl, M, y, 3.9, 0.4, LIME_XLT, pad=0.12, anchor=MIDDLE, paras=[(a_, dict(size=11.5, bold=True))])
    tb(sl, M + 4.05, y, 4.7, 0.4, b_, size=11, color=GREY, anchor=MIDDLE)
    line(sl, M + 8.75, y + 0.2, M + 9.05, y + 0.2, L700, 1.2, arrow=True)
    rect(sl, M + 9.1, y, CW - 9.1, 0.4, LIME_LT if 'Halqa' not in c_ and 'Counsel' not in c_ else C('F7F7F5'),
         pad=0.12, anchor=MIDDLE, paras=[(c_, dict(size=11))])
note(sl, 'Counsel to confirm each line in writing. Quotations and sections are verified in the Legal Position and the '
         'Bank Partnership Revisions.', y=6.65)

# ---- risks
sl = slide('Risks and Controls', 'Each risk placed by how likely and how damaging it is, with the control that '
                                'answers it. The placing is Halqa’s own assessment.')
gx, gy, gs = M + 0.6, 1.95, 1.45
for r_ in range(3):
    for c_ in range(3):
        sev = r_ + c_
        col = LIME_LT if sev <= 1 else (C('F6E7B4') if sev == 2 else C('F4C7C0'))
        rect(sl, gx + c_ * gs, gy + (2 - r_) * gs, gs - 0.05, gs - 0.05, col)
tb(sl, gx, gy + 3 * gs + 0.05, 3 * gs, 0.3, 'Likelihood', size=10.5, color=GREY, align=CEN)
tb(sl, M - 0.35, gy + 1.2 * gs, 0.9, 0.3, 'Impact', size=10.5, color=GREY)
risks = [('1', 2, 2), ('2', 1, 2), ('3', 1, 1), ('4', 0, 2), ('5', 1, 0), ('6', 2, 1)]
for n_, lk, im in risks:
    rect(sl, gx + lk * gs + 0.5, gy + (2 - im) * gs + 0.45, 0.42, 0.42, L700, pad=0, anchor=MIDDLE,
         paras=[(n_, dict(size=12, bold=True, color=WHITE, align=CEN))])
x0 = M + 5.3
riskt = [('1  The wallets take the category', 'JazzCash Committee reaches users first. Control: the bank partnership, '
                                             'a flat fee and credit records the wallet does not give.'),
         ('2  Defaults beyond the model', 'Control: bands, affordability, cover chosen by Mashreq, the hard stop.'),
         ('3  Money laundering or fraud', 'Control: Mashreq’s checks, payouts only to the member’s own account, fixed '
                                          'amounts on fixed dates.'),
         ('4  Rules change after a scandal', 'Control: work inside a bank’s licence, published rules for every '
                                             'circle.'),
         ('5  Technology failure', 'Control: outsourcing agreement, security tests, audit logs, daily '
                                   'reconciliation.'),
         ('6  Partnership stalls', 'Control: a pilot with fixed dates; Safepay routes keep collection alive.')]
for i, (t, d) in enumerate(riskt):
    y = 1.9 + i * 0.77
    tb(sl, x0, y, W - M - x0, 0.75, [(t, dict(size=12, bold=True, color=L700, gap=1)),
                                     (d, dict(size=10.5, spacing=1.05))])

# ---- technical
sl = slide('How the Systems Connect',
           'Three parties, seven steps. Halqa sends instructions; Mashreq moves the money and confirms each step.')
lanes = [('Member', M, icon_phone), ('Halqa platform', M + CW / 3, None), ('Mashreq systems', M + 2 * CW / 3,
                                                                          icon_bank)]
lw3 = CW / 3
for t, x, ic in lanes:
    rect(sl, x + 0.05, 1.85, lw3 - 0.1, 0.55, LIME if t == 'Halqa platform' else LIME_LT, pad=(0.7, 0.05, 0.1, 0.05),
         anchor=MIDDLE, paras=[(t, dict(size=14, bold=True))])
    if ic:
        ic(sl, x + 0.18, 1.9, 0.45, L700)
    else:
        sl.shapes.add_picture(MARK, E(x + 0.18), E(1.9), height=E(0.45))
    line(sl, x + lw3 / 2, 2.45, x + lw3 / 2, 6.55, RULE, 0.75, dash=True)
msgs = [(0, 2, 'Opens a Mashreq account inside the app'), (2, 1, 'Account and mandate references returned'),
        (1, 2, 'Collection list sent on the due date'), (2, 1, 'Signed confirmation for each payment'),
        (1, 0, 'Receipt, ledger and schedule updated'), (1, 2, 'Payout instruction: pot less arrears'),
        (2, 1, 'Daily statement for reconciliation')]
for i, (a_, b_, t) in enumerate(msgs):
    y = 2.75 + i * 0.52
    xa_ = M + a_ * lw3 + lw3 / 2
    xb_ = M + b_ * lw3 + lw3 / 2
    line(sl, xa_, y + 0.3, xb_, y + 0.3, L700, 1.4, arrow=True)
    tb(sl, min(xa_, xb_) + 0.1, y, abs(xb_ - xa_) - 0.2, 0.3, '%d  %s' % (i + 1, t), size=10.5, align=CEN)
note(sl, 'Halqa stores no card data, password or banking credential. Each party keeps an audit log of every '
         'instruction; breaks found at reconciliation are held and investigated.')

# ---- state
sl = slide('Where Halqa Stands', 'As of ' + DATE + '.')
state = [('Company', 'To be incorporated with the SECP; the chairman’s father named as director', AMBER),
         ('Product', 'Web application live since 20 July 2026, not open to the public; interface being rebuilt', LIME),
         ('Money', 'No real money has moved; payments run in a sandbox', AMBER),
         ('Partner bank', 'Mashreq Bank Pakistan proposed; a meeting expected around 10 October 2026', AMBER),
         ('Regulator', 'Informal guidance from the Chairman of the SECP; no formal filing', AMBER),
         ('Documents', 'Oraan Research, Bank Partnership Revisions, Partner Bank Proposition and a work register of '
                       '1,039 items', LIME)]
for i, (t, d, col) in enumerate(state):
    y = 1.9 + i * 0.72
    rect(sl, M, y + 0.12, 0.34, 0.34, col)
    tb(sl, M + 0.55, y + 0.05, 2.2, 0.5, t, size=14, bold=True, color=L700, anchor=MIDDLE)
    tb(sl, M + 2.8, y + 0.05, 6.4, 0.6, d, size=12, anchor=MIDDLE, spacing=1.05)
    line(sl, M, y + 0.66, M + 9.2, y + 0.66, RULE, 0.75)
x0 = M + 9.55
rect(sl, x0, 1.9, W - M - x0, 4.3, LIME_XLT, pad=0.2, paras=[
    ('Fix before real members', dict(size=14, font=DISPLAY, color=L700, gap=5)),
    ('Camera and location blocked by the permission policy', dict(size=10.5, bullet=True, gap=4, spacing=1.05)),
    ('Demonstration data to be wiped', dict(size=10.5, bullet=True, gap=4)),
    ('Error monitoring and versioned migrations', dict(size=10.5, bullet=True, gap=4, spacing=1.05)),
    ('Chat polling every six seconds', dict(size=10.5, bullet=True, gap=4)),
    ('Features of the August design removed', dict(size=10.5, bullet=True, spacing=1.05))])

# ---- approvals and pilot
sl = slide('Approvals and Pilot', 'Five approvals, then one six month cycle with about 1,000 members.')
tb(sl, M, 1.95, 5.6, 4.5, [
    'Mashreq’s board approves the partnership as an outsourcing arrangement, with Halqa assessed as a service '
    'provider.', 'Mashreq’s Shariah board approves the product for the Islamic window.',
    'Mashreq confirms whether the State Bank must approve the product or be notified, and any limits for a digital '
    'bank.', 'The partners sign a services agreement: data, security, service levels, audit rights, complaints and exit.',
    'A pilot runs within agreed limits before any public launch.'],
   size=12.5, num=True, gap=10, spacing=1.08, ind=0.32, bcolor=L700)
gx = 6.7
gl, mw = 1.95, 0.4
tb(sl, gx, 1.95, gl - 0.12, 0.25, 'Month', size=9, color=GREY, bold=True, align=R)
for m in range(10):
    tb(sl, gx + gl + m * mw, 1.95, mw, 0.25, str(m + 1), size=9, color=GREY, bold=True, align=CEN)
rows_ = [('Incorporation', 1, 1, L700), ('Agreement and approvals', 1, 2, L700), ('Integration and testing', 2, 3, LIME),
         ('Pilot cycle', 4, 9, LIME), ('Review', 6, 6, AMBER), ('Decision to scale', 10, 10, L700)]
for i, (lab, a_, b_, col) in enumerate(rows_):
    y = 2.3 + i * 0.42
    tb(sl, gx, y, gl - 0.12, 0.3, lab, size=10.5, align=R)
    rect(sl, gx + gl + (a_ - 1) * mw + 0.03, y + 0.03, (b_ - a_ + 1) * mw - 0.06, 0.26, col)
line(sl, gx + gl - 0.04, 2.25, gx + gl - 0.04, 4.8, RULE, 0.75)
tb(sl, gx, 5.0, W - M - gx, 1.6, [
    'Up to 100 monthly circles, about 1,000 members, on the Islamic window',
    'Every member opens a Mashreq account',
    'Measured: accounts opened, balances, on time payments, arrears recovered, complaints, cost of acquisition'],
   size=11.5, bullet=True, gap=5, spacing=1.05)

# ---- measures
sl = slide('Measures', 'Five numbers show whether this is working. A committee arrives as a formed group of ten to '
                       'fifteen people with a leader, so the measures follow the group.')
nums = [('Share of instalments collected through Mashreq', 'The health of collection'),
        ('Circles completed clean', 'The unit of value: a circle that never finishes proves nothing'),
        ('Circles per organiser', 'Whether growth compounds: a host’s second circle brings the whole group back'),
        ('Default after collecting, first 100 circles', 'When the modelled loss becomes a measured one'),
        ('Members who become organisers', 'The growth loop inside the product')]
for i, (t, d) in enumerate(nums):
    y = 2.0 + i * 0.86
    rect(sl, M, y, 0.62, 0.62, LIME, pad=0, anchor=MIDDLE, paras=[(str(i + 1), dict(size=20, font=DISPLAY,
                                                                                      align=CEN))])
    tb(sl, M + 0.85, y + 0.02, 6.0, 0.62, t, size=15, color=L700, font=DISPLAY, anchor=MIDDLE)
    tb(sl, M + 7.0, y + 0.02, W - M - 7.0, 0.62, d, size=12, anchor=MIDDLE, spacing=1.05)
    line(sl, M, y + 0.74, W - M, y + 0.74, RULE, 0.75)

# ---- open questions
sl = slide('Open Questions', 'What is still to be decided, and who decides it.')
qs = [('Mashreq', 'Which cover: its own guarantee, insurance or takaful'),
      ('Mashreq', 'Credit reporting through its own TASDEEQ membership, or to TASDEEQ directly'),
      ('Mashreq', 'Points as its stored value product, and the marketplace and rewards links'),
      ('Mashreq and the State Bank', 'Whether the product needs approval or notice, and any digital bank limits'),
      ('Counsel', 'Whether a turn sold for money is lending between members, and its Shariah standing'),
      ('The chairman', 'Point values: ten points a rupee earned, against one point a rupee of reward'),
      ('The chairman', 'The Hyper fee: Rs 15 or the current fee'),
      ('The chairman', 'The due date: the 8th for every circle, or set by payday')]
for i, (who, q) in enumerate(qs):
    y = 1.9 + i * 0.58
    rect(sl, M, y, 2.9, 0.48, LIME if who == 'The chairman' else LIME_LT, pad=0.12, anchor=MIDDLE,
         paras=[(who, dict(size=11.5, bold=True))])
    tb(sl, M + 3.1, y, CW - 3.1, 0.48, q, size=12.5, anchor=MIDDLE)

# ================================================================= SUMMARY ===
SECTION[0] = 'Summary'
sl = prs.slides.add_slide(BLANK)
PAGE[0] += 1
rect(sl, 9.3, 0, W - 9.3, H, LIME)
ring(sl, 9.3 + (W - 9.3) / 2, 3.75, 1.2, 0.17, 0.5)
sl.shapes.add_picture(LOGO, E(M), E(0.5), height=E(0.4))
tb(sl, M, 1.15, 8, 0.7, 'Summary', size=34, color=INK, font=DISPLAY)
tb(sl, M, 2.0, 8.3, 3.9, 'Committees are the savings habit Pakistan already has: paid monthly, on time, within trusted '
   'groups, mostly by women. Halqa runs the committees; Mashreq Bank Pakistan holds and moves the money under its '
   'licence, which answers the deposit question. Members get a flat fee, bands that protect every circle, '
   'cover chosen by the bank, a credit record from the first payment and the profit on their balances as points. '
   'Mashreq gets deposits, customers in groups, an Islamic product, overseas Pakistanis and a lending book built on '
   'real payment records. The next step is the meeting with Mashreq, then five approvals and a six month pilot with '
   'about 1,000 members.', size=15, color=INK, spacing=1.22)
tb(sl, M, 5.9, 8.3, 0.5, 'Halqa is the system. The bank is the machine.', size=22, color=L700, font=DISPLAY)
tb(sl, M, 6.55, 8, 0.3, 'Taha Amjed, Chairman   ' + DATE, size=11, color=INK, bold=True)

# ==================================================================== save ===
cp = prs.core_properties
cp.title = 'Halqa: the complete position'
cp.author = 'Halqa'
cp.last_modified_by = 'Halqa'
prs.save(OUT)

bad = []
for n, s in enumerate(prs.slides, 1):
    for shp in s.shapes:
        texts = []
        if shp.has_text_frame:
            texts.append(shp.text_frame.text)
        if getattr(shp, 'has_table', False) and shp.has_table:
            texts += [c.text_frame.text for r_ in shp.table.rows for c in r_.cells]
        for t in texts:
            for pat in (r'\byou\b', r'\byour\b', r'\bwe\b', r'\bour\b', u'[‒–—―]', r' - ',
                        r'(?i)\bjourney\b|\bunlock|\bseamless|\bempower|\bleverag|\brevolution'):
                if re.search(pat, t):
                    bad.append((n, pat, t[:70]))
print('slides', len(prs.slides), 'bytes', os.path.getsize(OUT))
for b in bad:
    print('CHECK', b)
