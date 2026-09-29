"""Design system for the Halqa bank decks: colours, type, grid and drawing helpers.

Rules carried from the chairman's deck rules: lime palette on white, Georgia titles
in lime 700, Arial body, no dashes, no pills, no accent lines under titles, sources
on every slide that carries a figure, page numbers.
"""
import os
import re

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")

# Colours
INK = "0C1408"
L700 = "41801A"
L600 = "56A323"
L500 = "6DC72A"
L300 = "A9DE7F"
T100 = "E3F8CA"
T50 = "F3FCE7"
GREY = "5B6456"
GREY2 = "8A9285"
LINE = "D5DCD1"
WHITE = "FFFFFF"
RED = "B3261E"

SERIF = "Georgia"
SANS = "Arial"

# Grid, inches
W, H = 13.333, 7.5
ML = 0.62
MR = 0.62
CW = W - ML - MR
TOP = 1.86
BOTTOM = 6.72

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}

BANNED = re.compile(r"[–—]|\b(journey|unlock|seamless|empower|leverage|"
                    r"revolution|truly|sakh|recorder)\b", re.I)


def rgb(hexstr):
    return RGBColor.from_string(hexstr)


def check_copy(text):
    """Stop the build on a dash or a banned word in slide or note text."""
    m = BANNED.search(text)
    if m:
        raise ValueError(f"Banned text {m.group(0)!r} in: {text[:80]}")


class Deck:
    def __init__(self, footer_name):
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.footer_name = footer_name
        self.count = 0

    def new(self, notes=""):
        s = self.prs.slides.add_slide(self.blank)
        self.count += 1
        if notes:
            check_copy(notes)
            s.notes_slide.notes_text_frame.text = notes
        return s

    def save(self, path):
        self.prs.save(path)


# Text

def _runs(p, text, size, color, bold, font, accent, italic=False):
    """Markup: **bold**, ^^accent colour^^, ~~grey~~."""
    parts = re.split(r"(\*\*|\^\^|~~)", text)
    state = {"**": False, "^^": False, "~~": False}
    for part in parts:
        if part in state:
            state[part] = not state[part]
            continue
        if not part:
            continue
        r = p.add_run()
        r.text = part
        f = r.font
        f.name = font
        f.size = Pt(size)
        f.bold = bold or state["**"]
        f.italic = italic
        c = color
        if state["^^"]:
            c = accent
        elif state["~~"]:
            c = GREY
        f.color.rgb = rgb(c)


def txt(slide, x, y, w, h, text, size=14, color=INK, bold=False, font=SANS,
        align="l", anchor="t", spacing=1.1, after=0, accent=L700, italic=False):
    """Text box. text is a string (\n separates paragraphs) or a list of strings."""
    paras = text if isinstance(text, list) else text.split("\n")
    for t in paras:
        check_copy(t)
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCHOR[anchor]
    for i, t in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        p.line_spacing = spacing
        if after:
            p.space_after = Pt(after)
        _runs(p, t, size, color, bold, font, accent, italic)
    return tb


def num(slide, x, y, w, value, caption, size=30, color=L700, cap_size=13, cap_h=0.6,
        align="l"):
    """A large figure in Georgia with a caption beneath it."""
    txt(slide, x, y, w, size / 60 + 0.1, value, size=size, color=color, font=SERIF,
        align=align)
    txt(slide, x, y + size / 60 + 0.12, w, cap_h, caption, size=cap_size, color=INK,
        align=align)


# Shapes

def _plain(shape):
    """Drop the theme style so no shadow or theme line is inherited."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)
    return shape


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    _plain(s)
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    s.text_frame.margin_left = s.text_frame.margin_right = 0
    return s


def oval(slide, cx, cy, d, fill=None, line=None, lw=1.25):
    return rect(slide, cx - d / 2, cy - d / 2, d, d, fill, line, lw, MSO_SHAPE.OVAL)


def pie(slide, cx, cy, d, start_deg, end_deg, fill):
    s = rect(slide, cx - d / 2, cy - d / 2, d, d, fill, None, 0, MSO_SHAPE.PIE)
    s.adjustments[0] = start_deg * 60000 / 100000
    s.adjustments[1] = end_deg * 60000 / 100000
    return s


def line(slide, x1, y1, x2, y2, color=INK, lw=0.75, arrow=False, dash=False,
         head=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    _plain(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", "dash")
    for tag, on in (("a:headEnd", head), ("a:tailEnd", arrow)):
        if on:
            e = etree.SubElement(ln, qn(tag))
            e.set("type", "triangle")
            e.set("w", "med")
            e.set("len", "med")
    return c


def hair(slide, x, y, w, color=LINE, lw=0.75):
    return line(slide, x, y, x + w, y, color, lw)


def picture(slide, name, x, y, w=None, h=None):
    path = os.path.join(ASSETS, name)
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)


# Page furniture

def logo(slide):
    picture(slide, "halqa_logo.png", W - MR - 1.12, 0.5, h=0.34)


def title(slide, heading, statement=None, sub=None):
    """Noun heading in Georgia, then one plain statement: the slide's point."""
    txt(slide, ML, 0.4, 10.4, 0.62, heading, size=28, color=L700, font=SERIF)
    if statement:
        txt(slide, ML, 1.08, CW, 0.45, statement, size=18, color=INK)
    if sub:
        txt(slide, ML, 1.46, CW, 0.3, sub, size=13, color=GREY)
    logo(slide)


def source(slide, text):
    txt(slide, ML, 6.9, 11.1, 0.4, text, size=9, color=GREY, spacing=1.0)


def page(slide, n):
    txt(slide, W - MR - 0.6, 6.95, 0.6, 0.25, str(n), size=10, color=GREY, align="r")


def content_slide(deck, heading, statement, notes, src=None, sub=None):
    s = deck.new(notes)
    title(s, heading, statement, sub)
    if src:
        source(s, src)
    page(s, deck.count)
    return s


def emu(inches):
    return Emu(int(inches * 914400))
